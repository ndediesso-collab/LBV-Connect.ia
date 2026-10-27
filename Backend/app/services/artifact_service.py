from __future__ import annotations

import ast
import csv
import hashlib
import io
import json
import mimetypes
import re
import sqlite3
import zipfile

from dataclasses import dataclass, field
from html import escape as html_escape
from pathlib import PurePosixPath
from typing import Any, Iterable
from uuid import uuid4


class ArtifactError(RuntimeError):
    """Erreur générale de génération de fichier Oria."""


class ArtifactDependencyError(ArtifactError):
    """Une dépendance optionnelle nécessaire à un format est absente."""


class UnsupportedArtifactTypeError(ArtifactError):
    """Le format demandé n'est pas générable nativement."""


@dataclass(slots=True)
class ArtifactResult:
    """
    Résultat binaire produit par ArtifactService.

    `data` reste en mémoire. La persistance (Supabase Storage / autre stockage)
    sera gérée par une couche dédiée afin que ce service reste centré sur
    la fabrication du fichier.
    """

    artifact_id: str
    filename: str
    extension: str
    mime_type: str
    category: str
    data: bytes
    size_bytes: int
    sha256: str
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(
        self,
        *,
        include_data: bool = False,
    ) -> dict[str, Any]:
        payload = {
            "artifact_id": self.artifact_id,
            "filename": self.filename,
            "extension": self.extension,
            "mime_type": self.mime_type,
            "category": self.category,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
            "metadata": dict(self.metadata),
        }

        if include_data:
            payload["data"] = self.data

        return payload


class ArtifactService:
    """
    Génération native de fichiers pour Oria.

    Objectif :
    ---------
    Recevoir le contenu produit par le modèle puis fabriquer un véritable
    fichier téléchargeable sans demander au frontend de reconstruire lui-même
    les formats binaires.

    Ce service couvre :
    - documents : PDF, DOCX, RTF, ODT, EPUB ;
    - tableurs / données : XLSX, ODS, CSV, TSV, JSON, JSONL, SQLite ;
    - présentations : PPTX, ODP ;
    - Web / texte / code : TXT, MD, HTML, CSS, JS, TS, TSX, JSX, PY, SQL,
      XML, YAML, YML, TOML, SVG, TEX, ICS, VCF et de nombreuses extensions
      textuelles ;
    - archives : ZIP.

    Les images, vidéos et futurs fichiers audio générés par des modèles dédiés
    restent dans leurs services média respectifs. Ils pourront néanmoins être
    placés dans une archive ZIP par ce service.

    Pour les extensions textuelles non répertoriées, Oria peut utiliser
    `allow_text_fallback=True` : le contenu est écrit tel quel en UTF-8.
    Cela permet de générer pratiquement tout fichier source/configuration
    textuel sans multiplier les générateurs spécialisés.
    """

    MAX_ARTIFACT_BYTES = 50 * 1024 * 1024
    MAX_ARCHIVE_FILES = 100
    # Empêche qu'un petit ZIP fortement compressé contienne des centaines
    # de Mo de fichiers source en mémoire. Indépendant de la taille du ZIP.
    MAX_ARCHIVE_UNCOMPRESSED_BYTES = 4 * MAX_ARTIFACT_BYTES
    MAX_ARCHIVE_DEPTH = 2
    MAX_FILENAME_LENGTH = 140

    DEFAULT_TEXT_ENCODING = "utf-8"

    TYPE_ALIASES = {
        # Noms de langages courants acceptés en plus des extensions.
        "python": "py",
        "javascript": "js",
        "typescript": "ts",
        "bash": "sh",
        "c++": "cpp",
        "csharp": "cs",
        "golang": "go",
        "rust": "rs",
        "ruby": "rb",
        "shell": "sh",
        "powershell": "ps1",
        "text": "txt",
        "plaintext": "txt",
        "markdown": "md",
        "word": "docx",
        "doc": "docx",
        "excel": "xlsx",
        "spreadsheet": "xlsx",
        "powerpoint": "pptx",
        "slides": "pptx",
        "presentation": "pptx",
        "database": "sqlite",
        "db": "sqlite",
        "sqlite3": "sqlite",
        "archive": "zip",
        "latex": "tex",
        "calendar": "ics",
        "contact": "vcf",
        "ebook": "epub",
        "open-document-text": "odt",
        "open-document-sheet": "ods",
        "open-document-presentation": "odp",
    }

    MIME_TYPES = {
        "txt": "text/plain; charset=utf-8",
        "md": "text/markdown; charset=utf-8",
        "html": "text/html; charset=utf-8",
        "htm": "text/html; charset=utf-8",
        "css": "text/css; charset=utf-8",
        "js": "text/javascript; charset=utf-8",
        "mjs": "text/javascript; charset=utf-8",
        "cjs": "text/javascript; charset=utf-8",
        "ts": "text/typescript; charset=utf-8",
        "tsx": "text/typescript; charset=utf-8",
        "jsx": "text/javascript; charset=utf-8",
        "py": "text/x-python; charset=utf-8",
        "java": "text/x-java-source; charset=utf-8",
        "c": "text/x-c; charset=utf-8",
        "h": "text/x-c; charset=utf-8",
        "cpp": "text/x-c++; charset=utf-8",
        "hpp": "text/x-c++; charset=utf-8",
        "cs": "text/plain; charset=utf-8",
        "go": "text/plain; charset=utf-8",
        "rs": "text/plain; charset=utf-8",
        "php": "application/x-httpd-php; charset=utf-8",
        "rb": "text/plain; charset=utf-8",
        "sh": "text/x-shellscript; charset=utf-8",
        "ps1": "text/plain; charset=utf-8",
        "bat": "text/plain; charset=utf-8",
        "sql": "application/sql; charset=utf-8",
        "xml": "application/xml; charset=utf-8",
        "yaml": "application/yaml; charset=utf-8",
        "yml": "application/yaml; charset=utf-8",
        "toml": "application/toml; charset=utf-8",
        "ini": "text/plain; charset=utf-8",
        "env": "text/plain; charset=utf-8",
        "log": "text/plain; charset=utf-8",
        "tex": "application/x-tex; charset=utf-8",
        "rtf": "application/rtf",
        "csv": "text/csv; charset=utf-8",
        "tsv": "text/tab-separated-values; charset=utf-8",
        "json": "application/json; charset=utf-8",
        "jsonl": "application/x-ndjson; charset=utf-8",
        "ndjson": "application/x-ndjson; charset=utf-8",
        "svg": "image/svg+xml; charset=utf-8",
        "ics": "text/calendar; charset=utf-8",
        "vcf": "text/vcard; charset=utf-8",
        "pdf": "application/pdf",
        "docx": (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        "xlsx": (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        "pptx": (
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
        "odt": "application/vnd.oasis.opendocument.text",
        "ods": "application/vnd.oasis.opendocument.spreadsheet",
        "odp": "application/vnd.oasis.opendocument.presentation",
        "epub": "application/epub+zip",
        "sqlite": "application/vnd.sqlite3",
        "zip": "application/zip",
    }

    CATEGORY_BY_TYPE = {
        "pdf": "document",
        "docx": "document",
        "rtf": "document",
        "odt": "document",
        "epub": "document",
        "xlsx": "spreadsheet",
        "ods": "spreadsheet",
        "csv": "data",
        "tsv": "data",
        "json": "data",
        "jsonl": "data",
        "ndjson": "data",
        "sqlite": "database",
        "pptx": "presentation",
        "odp": "presentation",
        "zip": "archive",
        "svg": "image",
        "ics": "calendar",
        "vcf": "contact",
    }

    KNOWN_TEXT_TYPES = {
        "txt", "md", "html", "htm", "css", "js", "mjs", "cjs",
        "ts", "tsx", "jsx", "py", "java", "c", "h", "cpp", "hpp",
        "cs", "go", "rs", "php", "rb", "sh", "ps1", "bat", "sql",
        "xml", "yaml", "yml", "toml", "ini", "env", "log", "tex",
        "svg", "ics", "vcf",
    }

    # Sous-ensemble des formats textuels à traiter comme des sources. Ne
    # change pas leur catégorie publique (text), pour rester compatible avec
    # les frontends et les repositories déjà déployés.
    CODE_SOURCE_TYPES = frozenset({
        "py", "js", "mjs", "cjs", "ts", "tsx", "jsx", "html", "htm",
        "css", "sql", "xml", "yaml", "yml", "toml", "svg", "sh",
        "ps1", "bat", "java", "c", "h", "cpp", "hpp", "cs",
        "go", "rs", "php", "rb", "ini", "env",
    })

    BINARY_TYPES = {
        "pdf", "docx", "xlsx", "pptx", "odt", "ods", "odp",
        "epub", "sqlite", "zip",
    }

    def supported_types(self) -> dict[str, list[str]]:
        """Retourne les formats natifs exposables au frontend."""
        return {
            "documents": ["pdf", "docx", "rtf", "odt", "epub"],
            "spreadsheets": ["xlsx", "ods", "csv", "tsv"],
            "presentations": ["pptx", "odp"],
            "data": ["json", "jsonl", "ndjson", "sqlite"],
            "web_and_code": sorted(self.KNOWN_TEXT_TYPES),
            "archives": ["zip"],
        }

    def create(
        self,
        *,
        file_type: str,
        content: Any,
        filename: str | None = None,
        title: str | None = None,
        metadata: dict[str, Any] | None = None,
        allow_text_fallback: bool = True,
        _archive_depth: int = 0,
    ) -> ArtifactResult:
        """
        Fabrique un fichier en mémoire.

        `content` dépend du format :
        - documents : str ou structure de blocs ;
        - tableurs : lignes, liste de dictionnaires, ou {"sheets": [...]};
        - présentations : {"slides": [{"title": ..., "bullets": [...]}]};
        - ZIP : {"files": [{"name": ..., "file_type": ..., "content": ...}]};
        - SQLite : {"tables": [{"name": ..., "rows": [...]}]}.
        """
        normalized_type = self._normalize_type(file_type, filename)
        safe_filename = self._resolve_filename(
            filename=filename,
            title=title,
            extension=normalized_type,
        )

        if normalized_type == "pdf":
            data = self._create_pdf(content, title=title)
        elif normalized_type == "docx":
            data = self._create_docx(content, title=title)
        elif normalized_type == "xlsx":
            data = self._create_xlsx(content)
        elif normalized_type == "pptx":
            data = self._create_pptx(content, title=title)
        elif normalized_type == "odt":
            data = self._create_odt(content, title=title)
        elif normalized_type == "ods":
            data = self._create_ods(content)
        elif normalized_type == "odp":
            data = self._create_odp(content, title=title)
        elif normalized_type == "epub":
            data = self._create_epub(content, title=title)
        elif normalized_type == "sqlite":
            data = self._create_sqlite(content)
        elif normalized_type == "zip":
            data = self._create_zip(
                content,
                depth=_archive_depth,
                # Un ZIP demandé depuis le Chat peut contenir un vrai projet.
                # Contrairement à une archive de données, aucun chemin ni
                # nom de fichier ne doit être modifié silencieusement.
                strict_project=(metadata or {}).get("source") == "oria-chat",
            )
        elif normalized_type == "rtf":
            data = self._create_rtf(content)
        elif normalized_type == "csv":
            data = self._create_delimited(content, delimiter=",")
        elif normalized_type == "tsv":
            data = self._create_delimited(content, delimiter="\t")
        elif normalized_type == "json":
            data = self._create_json(content)
        elif normalized_type in {"jsonl", "ndjson"}:
            data = self._create_jsonl(content)
        elif normalized_type in self.KNOWN_TEXT_TYPES:
            # ai.py fournit déjà un bloc de source extrait du Markdown.
            # Si une réponse Oria est accidentellement un dictionnaire/liste,
            # ne pas fabriquer un fichier .py/.js contenant du JSON trompeur.
            if (
                normalized_type in self.CODE_SOURCE_TYPES
                and (metadata or {}).get("source") == "oria-chat"
                and not isinstance(content, (str, bytes))
            ):
                raise ArtifactError(
                    "Un fichier source ORIA doit contenir du code texte "
                    "ou des octets, pas un objet structuré."
                )
            data = self._create_text(content)
            if (
                normalized_type in self.CODE_SOURCE_TYPES
                and (metadata or {}).get("source") == "oria-chat"
            ):
                self._validate_chat_source(
                    filename=safe_filename,
                    extension=normalized_type,
                    data=data,
                )
        elif allow_text_fallback:
            if self._looks_binary_extension(normalized_type):
                raise UnsupportedArtifactTypeError(
                    f"Le format binaire '.{normalized_type}' n'a pas encore "
                    "de générateur natif Oria."
                )
            data = self._create_text(content)
        else:
            raise UnsupportedArtifactTypeError(
                f"Format de fichier non supporté : {normalized_type}"
            )

        if not isinstance(data, bytes):
            raise ArtifactError(
                "Le générateur de fichier n'a pas retourné des octets."
            )

        self._validate_size(data)

        mime_type = self._resolve_mime_type(normalized_type)
        category = self.CATEGORY_BY_TYPE.get(
            normalized_type,
            "text" if normalized_type not in self.BINARY_TYPES else "file",
        )

        artifact_id = str(uuid4())
        digest = hashlib.sha256(data).hexdigest()

        result_metadata = dict(metadata or {})
        result_metadata.setdefault("title", title or "")
        result_metadata["generator"] = "oria-native-artifact-service"
        result_metadata["format"] = normalized_type

        return ArtifactResult(
            artifact_id=artifact_id,
            filename=safe_filename,
            extension=normalized_type,
            mime_type=mime_type,
            category=category,
            data=data,
            size_bytes=len(data),
            sha256=digest,
            metadata=result_metadata,
        )

    # ============================================================
    # NORMALISATION / SÉCURITÉ
    # ============================================================

    def _normalize_type(
        self,
        file_type: str,
        filename: str | None,
    ) -> str:
        value = str(file_type or "").strip().lower()
        value = value.lstrip(".")

        if not value and filename:
            suffix = PurePosixPath(filename).suffix
            value = suffix.lower().lstrip(".")

        value = self.TYPE_ALIASES.get(value, value)

        if not value:
            raise UnsupportedArtifactTypeError(
                "Le type de fichier est requis."
            )

        if re.fullmatch(r"[a-z0-9][a-z0-9+._-]{0,31}", value) is None:
            raise UnsupportedArtifactTypeError(
                "Type de fichier invalide."
            )

        return value

    def _resolve_filename(
        self,
        *,
        filename: str | None,
        title: str | None,
        extension: str,
    ) -> str:
        base = str(filename or title or "document-oria").strip()
        base = PurePosixPath(base.replace("\\", "/")).name
        # Un fichier .env est un nom de configuration réel, pas env.env.
        if filename and base == ".env" and extension == "env":
            return base
        base = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "-", base)
        base = re.sub(r"\s+", " ", base).strip(" .-_")

        if not base:
            base = "document-oria"

        expected_suffix = f".{extension.lower()}"

        if not base.lower().endswith(expected_suffix):
            # Retirer une extension existante seulement si elle ressemble
            # réellement à une extension de fichier.
            path = PurePosixPath(base)
            if path.suffix and len(path.suffix) <= 12:
                base = path.stem
            base = f"{base}{expected_suffix}"

        if len(base) > self.MAX_FILENAME_LENGTH:
            suffix = f".{extension}"
            stem_limit = max(
                1,
                self.MAX_FILENAME_LENGTH - len(suffix),
            )
            stem = PurePosixPath(base).stem[:stem_limit].rstrip(" .-_")
            base = f"{stem or 'document-oria'}{suffix}"

        return base

    def _resolve_mime_type(
        self,
        extension: str,
    ) -> str:
        if extension in self.MIME_TYPES:
            return self.MIME_TYPES[extension]

        guessed, _ = mimetypes.guess_type(f"file.{extension}")

        if guessed:
            if guessed.startswith("text/"):
                return f"{guessed}; charset=utf-8"
            return guessed

        return "application/octet-stream"

    def _validate_size(
        self,
        data: bytes,
    ) -> None:
        if not data:
            raise ArtifactError("Le fichier généré est vide.")

        if len(data) > self.MAX_ARTIFACT_BYTES:
            raise ArtifactError(
                "Le fichier généré dépasse la taille maximale autorisée "
                f"({self.MAX_ARTIFACT_BYTES // (1024 * 1024)} Mo)."
            )

    @staticmethod
    def _looks_binary_extension(
        extension: str,
    ) -> bool:
        binary_extensions = {
            "exe", "dll", "so", "dylib", "bin", "iso", "dmg", "apk",
            "aab", "msi", "deb", "rpm", "psd", "ai", "sketch", "blend",
            "dwg", "dxf", "step", "stp", "fbx", "obj", "mp3", "wav",
            "aac", "flac", "ogg", "mp4", "mov", "avi", "mkv", "webm",
            "png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff",
        }
        return extension in binary_extensions

    # ============================================================
    # TEXTE / DONNÉES SIMPLES
    # ============================================================

    def _create_text(
        self,
        content: Any,
    ) -> bytes:
        if isinstance(content, bytes):
            return content

        if isinstance(content, str):
            text = content
        elif isinstance(content, (dict, list, tuple)):
            text = json.dumps(
                content,
                ensure_ascii=False,
                indent=2,
                default=str,
            )
        elif content is None:
            text = ""
        else:
            text = str(content)

        return text.encode(self.DEFAULT_TEXT_ENCODING)

    def _create_json(
        self,
        content: Any,
    ) -> bytes:
        if isinstance(content, str):
            stripped = content.strip()
            try:
                parsed = json.loads(stripped)
            except json.JSONDecodeError:
                parsed = content
        else:
            parsed = content

        return (
            json.dumps(
                parsed,
                ensure_ascii=False,
                indent=2,
                default=str,
            )
            + "\n"
        ).encode(self.DEFAULT_TEXT_ENCODING)

    def _create_jsonl(
        self,
        content: Any,
    ) -> bytes:
        if isinstance(content, str):
            return content.rstrip().encode(self.DEFAULT_TEXT_ENCODING) + b"\n"

        rows = content
        if isinstance(content, dict):
            rows = content.get("rows", content.get("items", [content]))

        if not isinstance(rows, (list, tuple)):
            rows = [rows]

        output = "\n".join(
            json.dumps(
                row,
                ensure_ascii=False,
                default=str,
            )
            for row in rows
        )

        return (output + "\n").encode(self.DEFAULT_TEXT_ENCODING)

    def _create_delimited(
        self,
        content: Any,
        *,
        delimiter: str,
    ) -> bytes:
        rows = self._normalize_rows(content)

        buffer = io.StringIO(newline="")
        writer = csv.writer(
            buffer,
            delimiter=delimiter,
            lineterminator="\n",
        )

        for row in rows:
            writer.writerow(
                "" if value is None else value
                for value in row
            )

        # UTF-8 BOM améliore l'ouverture directe dans Excel sur Windows.
        return (
            "\ufeff" + buffer.getvalue()
        ).encode(self.DEFAULT_TEXT_ENCODING)

    def _create_rtf(
        self,
        content: Any,
    ) -> bytes:
        text = self._document_plain_text(content)

        def rtf_escape(value: str) -> str:
            parts: list[str] = []
            for char in value:
                code = ord(char)
                if char in "\\{}":
                    parts.append("\\" + char)
                elif char == "\n":
                    parts.append("\\par\n")
                elif code > 127:
                    signed = code if code <= 32767 else code - 65536
                    parts.append(f"\\u{signed}?")
                else:
                    parts.append(char)
            return "".join(parts)

        body = rtf_escape(text)
        rtf = (
            r"{\rtf1\ansi\deff0"
            "\n"
            r"{\fonttbl{\f0 Arial;}}"
            "\n"
            r"\fs22 "
            + body
            + "\n}"
        )
        return rtf.encode("ascii", errors="strict")

    # ============================================================
    # DOCUMENTS
    # ============================================================

    def _create_pdf(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> bytes:
        try:
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                ListFlowable,
                ListItem,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
            from reportlab.lib import colors
        except ImportError as error:
            raise ArtifactDependencyError(
                "PDF natif : installez 'reportlab'."
            ) from error

        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
            title=title or "Document Oria",
            author="Oria",
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "OriaTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            spaceAfter=14,
        )
        body_style = styles["BodyText"]
        body_style.leading = 15
        heading1 = styles["Heading1"]
        heading2 = styles["Heading2"]
        heading3 = styles["Heading3"]

        story: list[Any] = []

        if title:
            story.append(
                Paragraph(
                    html_escape(title),
                    title_style,
                )
            )

        for block in self._normalize_document_blocks(content):
            kind = block["type"]
            value = block.get("text", "")

            if kind == "h1":
                story.append(Paragraph(html_escape(value), heading1))
            elif kind == "h2":
                story.append(Paragraph(html_escape(value), heading2))
            elif kind == "h3":
                story.append(Paragraph(html_escape(value), heading3))
            elif kind == "bullet":
                story.append(
                    ListFlowable(
                        [
                            ListItem(
                                Paragraph(
                                    html_escape(item),
                                    body_style,
                                )
                            )
                            for item in block.get("items", [])
                        ],
                        bulletType="bullet",
                    )
                )
            elif kind == "table":
                rows = block.get("rows", [])
                if rows:
                    table = Table(rows, repeatRows=1)
                    table.setStyle(
                        TableStyle(
                            [
                                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                                ("TOPPADDING", (0, 0), (-1, -1), 4),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                            ]
                        )
                    )
                    story.append(table)
            elif kind == "spacer":
                story.append(Spacer(1, 8))
            else:
                safe = html_escape(value).replace("\n", "<br/>")
                story.append(Paragraph(safe or " ", body_style))

            story.append(Spacer(1, 5))

        document.build(story)
        return buffer.getvalue()

    def _create_docx(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> bytes:
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError as error:
            raise ArtifactDependencyError(
                "DOCX natif : installez 'python-docx'."
            ) from error

        document = Document()

        styles = document.styles
        if "Normal" in styles:
            styles["Normal"].font.name = "Arial"
            styles["Normal"].font.size = Pt(11)

        if title:
            document.add_heading(title, level=0)

        for block in self._normalize_document_blocks(content):
            kind = block["type"]

            if kind == "h1":
                document.add_heading(block.get("text", ""), level=1)
            elif kind == "h2":
                document.add_heading(block.get("text", ""), level=2)
            elif kind == "h3":
                document.add_heading(block.get("text", ""), level=3)
            elif kind == "bullet":
                for item in block.get("items", []):
                    document.add_paragraph(
                        str(item),
                        style="List Bullet",
                    )
            elif kind == "table":
                rows = block.get("rows", [])
                if rows:
                    cols = max(len(row) for row in rows)
                    table = document.add_table(
                        rows=len(rows),
                        cols=max(cols, 1),
                    )
                    table.style = "Table Grid"
                    for row_index, row in enumerate(rows):
                        for col_index, value in enumerate(row):
                            table.cell(
                                row_index,
                                col_index,
                            ).text = "" if value is None else str(value)
            elif kind == "spacer":
                document.add_paragraph("")
            else:
                document.add_paragraph(block.get("text", ""))

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _create_odt(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> bytes:
        try:
            from odf.opendocument import OpenDocumentText
            from odf.text import H, P
        except ImportError as error:
            raise ArtifactDependencyError(
                "ODT natif : installez 'odfpy'."
            ) from error

        document = OpenDocumentText()

        if title:
            document.text.addElement(H(outlinelevel=1, text=title))

        for block in self._normalize_document_blocks(content):
            kind = block["type"]
            if kind in {"h1", "h2", "h3"}:
                level = int(kind[-1])
                document.text.addElement(
                    H(
                        outlinelevel=level,
                        text=block.get("text", ""),
                    )
                )
            elif kind == "bullet":
                for item in block.get("items", []):
                    document.text.addElement(P(text=f"• {item}"))
            elif kind == "table":
                for row in block.get("rows", []):
                    document.text.addElement(
                        P(
                            text=" | ".join(
                                "" if cell is None else str(cell)
                                for cell in row
                            )
                        )
                    )
            else:
                document.text.addElement(
                    P(text=block.get("text", ""))
                )

        buffer = io.BytesIO()
        document.write(buffer)
        return buffer.getvalue()

    def _create_epub(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> bytes:
        """
        Génère un EPUB 3 minimal sans dépendance externe.

        EPUB est un conteneur ZIP normalisé. Le générer directement permet à
        Oria d'éviter une dépendance supplémentaire uniquement pour ce format.
        """
        book_id = str(uuid4())
        book_title = title or "Document Oria"
        chapters = self._normalize_epub_chapters(content)

        container_xml = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0"
 xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="EPUB/content.opf"
     media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""

        manifest_items = []
        spine_items = []
        nav_items = []
        chapter_files: list[tuple[str, str]] = []

        for index, chapter in enumerate(chapters, start=1):
            chapter_id = f"chapter-{index}"
            filename = f"chapter_{index}.xhtml"
            chapter_title = html_escape(chapter["title"])

            paragraphs = "\n".join(
                f"<p>{html_escape(paragraph)}</p>"
                for paragraph in chapter["paragraphs"]
            )

            chapter_xhtml = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" lang="fr">
<head>
  <meta charset="utf-8"/>
  <title>{chapter_title}</title>
</head>
<body>
  <h1>{chapter_title}</h1>
  {paragraphs}
</body>
</html>
"""

            chapter_files.append((filename, chapter_xhtml))
            manifest_items.append(
                f'    <item id="{chapter_id}" href="{filename}" '
                'media-type="application/xhtml+xml"/>'
            )
            spine_items.append(
                f'    <itemref idref="{chapter_id}"/>'
            )
            nav_items.append(
                f'      <li><a href="{filename}">{chapter_title}</a></li>'
            )

        nav_xhtml = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml"
 xmlns:epub="http://www.idpf.org/2007/ops" lang="fr">
<head>
  <meta charset="utf-8"/>
  <title>Sommaire</title>
</head>
<body>
  <nav epub:type="toc" id="toc">
    <h1>Sommaire</h1>
    <ol>
{chr(10).join(nav_items)}
    </ol>
  </nav>
</body>
</html>
"""

        content_opf = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf"
 version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">{book_id}</dc:identifier>
    <dc:title>{html_escape(book_title)}</dc:title>
    <dc:language>fr</dc:language>
    <dc:creator>Oria</dc:creator>
  </metadata>
  <manifest>
    <item id="nav" href="nav.xhtml"
     media-type="application/xhtml+xml" properties="nav"/>
{chr(10).join(manifest_items)}
  </manifest>
  <spine>
{chr(10).join(spine_items)}
  </spine>
</package>
"""

        buffer = io.BytesIO()

        with zipfile.ZipFile(buffer, "w") as archive:
            # EPUB exige que mimetype soit le premier fichier et non compressé.
            archive.writestr(
                "mimetype",
                "application/epub+zip",
                compress_type=zipfile.ZIP_STORED,
            )
            archive.writestr(
                "META-INF/container.xml",
                container_xml,
                compress_type=zipfile.ZIP_DEFLATED,
            )
            archive.writestr(
                "EPUB/content.opf",
                content_opf,
                compress_type=zipfile.ZIP_DEFLATED,
            )
            archive.writestr(
                "EPUB/nav.xhtml",
                nav_xhtml,
                compress_type=zipfile.ZIP_DEFLATED,
            )

            for filename, chapter_xhtml in chapter_files:
                archive.writestr(
                    f"EPUB/{filename}",
                    chapter_xhtml,
                    compress_type=zipfile.ZIP_DEFLATED,
                )

        return buffer.getvalue()

    # ============================================================
    # TABLEURS
    # ============================================================

    def _create_xlsx(
        self,
        content: Any,
    ) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as error:
            raise ArtifactDependencyError(
                "XLSX natif : installez 'openpyxl'."
            ) from error

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        sheets = self._normalize_sheets(content)

        for sheet_index, sheet_spec in enumerate(sheets, start=1):
            sheet_name = self._safe_sheet_name(
                sheet_spec.get("name") or f"Feuille {sheet_index}"
            )
            worksheet = workbook.create_sheet(sheet_name)

            rows = sheet_spec.get("rows", [])

            for row_index, row in enumerate(rows, start=1):
                for col_index, value in enumerate(row, start=1):
                    worksheet.cell(
                        row=row_index,
                        column=col_index,
                        value=value,
                    )

            if rows:
                for cell in worksheet[1]:
                    cell.font = Font(bold=True)

            # Largeur approximative utile sans calcul coûteux.
            max_cols = max((len(row) for row in rows), default=0)
            for col_index in range(1, min(max_cols, 50) + 1):
                letter = worksheet.cell(
                    row=1,
                    column=col_index,
                ).column_letter
                width = 12
                for row in rows[:200]:
                    if col_index - 1 < len(row):
                        width = min(
                            45,
                            max(
                                width,
                                len(str(row[col_index - 1] or "")) + 2,
                            ),
                        )
                worksheet.column_dimensions[letter].width = width

        if not workbook.sheetnames:
            workbook.create_sheet("Feuille 1")

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _create_ods(
        self,
        content: Any,
    ) -> bytes:
        try:
            from odf.opendocument import OpenDocumentSpreadsheet
            from odf.table import Table, TableCell, TableRow
            from odf.text import P
        except ImportError as error:
            raise ArtifactDependencyError(
                "ODS natif : installez 'odfpy'."
            ) from error

        document = OpenDocumentSpreadsheet()

        for index, sheet_spec in enumerate(
            self._normalize_sheets(content),
            start=1,
        ):
            table = Table(
                name=self._safe_sheet_name(
                    sheet_spec.get("name") or f"Feuille {index}"
                )
            )

            for row in sheet_spec.get("rows", []):
                table_row = TableRow()

                for value in row:
                    cell = TableCell(
                        valuetype="string",
                    )
                    cell.addElement(
                        P(text="" if value is None else str(value))
                    )
                    table_row.addElement(cell)

                table.addElement(table_row)

            document.spreadsheet.addElement(table)

        buffer = io.BytesIO()
        document.write(buffer)
        return buffer.getvalue()

    # ============================================================
    # PRÉSENTATIONS
    # ============================================================

    def _create_pptx(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> bytes:
        try:
            from pptx import Presentation
        except ImportError as error:
            raise ArtifactDependencyError(
                "PPTX natif : installez 'python-pptx'."
            ) from error

        presentation = Presentation()
        slides = self._normalize_slides(content, title=title)

        # Supprimer la présentation vide implicite n'est pas nécessaire :
        # python-pptx démarre sans slides.
        for slide_spec in slides:
            layout = presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)

            slide.shapes.title.text = slide_spec["title"]

            body = slide.placeholders[1]
            text_frame = body.text_frame
            text_frame.clear()

            bullets = slide_spec.get("bullets", [])

            if not bullets:
                bullets = [""]

            for index, bullet in enumerate(bullets):
                paragraph = (
                    text_frame.paragraphs[0]
                    if index == 0
                    else text_frame.add_paragraph()
                )
                paragraph.text = str(bullet)
                paragraph.level = 0

        buffer = io.BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()

    def _create_odp(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> bytes:
        try:
            from odf.draw import Frame, Page, TextBox
            from odf.opendocument import OpenDocumentPresentation
            from odf.text import H, P
        except ImportError as error:
            raise ArtifactDependencyError(
                "ODP natif : installez 'odfpy'."
            ) from error

        document = OpenDocumentPresentation()
        slides = self._normalize_slides(content, title=title)

        for index, slide_spec in enumerate(slides, start=1):
            page = Page(
                name=f"slide{index}",
                masterpagename="Default",
            )

            title_frame = Frame(
                width="24cm",
                height="2cm",
                x="1cm",
                y="1cm",
            )
            title_box = TextBox()
            title_box.addElement(
                H(
                    outlinelevel=1,
                    text=slide_spec["title"],
                )
            )
            title_frame.addElement(title_box)
            page.addElement(title_frame)

            body_frame = Frame(
                width="24cm",
                height="14cm",
                x="1cm",
                y="3.5cm",
            )
            body_box = TextBox()
            for bullet in slide_spec.get("bullets", []):
                body_box.addElement(P(text=f"• {bullet}"))
            body_frame.addElement(body_box)
            page.addElement(body_frame)

            document.presentation.addElement(page)

        buffer = io.BytesIO()
        document.write(buffer)
        return buffer.getvalue()

    # ============================================================
    # SQLITE
    # ============================================================

    def _create_sqlite(
        self,
        content: Any,
    ) -> bytes:
        tables = self._normalize_database_tables(content)

        # sqlite3 n'écrit pas directement dans BytesIO. On utilise donc
        # serialize(), disponible sur les versions modernes de Python.
        connection = sqlite3.connect(":memory:")

        try:
            for table_spec in tables:
                table_name = self._safe_sql_identifier(
                    table_spec["name"]
                )
                rows = table_spec["rows"]

                if not rows:
                    connection.execute(
                        f'CREATE TABLE "{table_name}" '
                        '("_empty" TEXT)'
                    )
                    continue

                headers = list(rows[0].keys())

                safe_headers = [
                    self._safe_sql_identifier(header)
                    for header in headers
                ]

                columns_sql = ", ".join(
                    f'"{header}" TEXT'
                    for header in safe_headers
                )

                connection.execute(
                    f'CREATE TABLE "{table_name}" ({columns_sql})'
                )

                placeholders = ", ".join("?" for _ in safe_headers)
                insert_sql = (
                    f'INSERT INTO "{table_name}" VALUES '
                    f"({placeholders})"
                )

                connection.executemany(
                    insert_sql,
                    [
                        [
                            None if row.get(header) is None
                            else str(row.get(header))
                            for header in headers
                        ]
                        for row in rows
                    ],
                )

            connection.commit()

            serialize = getattr(connection, "serialize", None)
            if not callable(serialize):
                raise ArtifactError(
                    "Cette version de Python ne permet pas de sérialiser "
                    "une base SQLite en mémoire."
                )

            return bytes(serialize())

        finally:
            connection.close()

    # ============================================================
    # ZIP
    # ============================================================

    @staticmethod
    def _validate_chat_source(
        *,
        filename: str,
        extension: str,
        data: bytes,
    ) -> None:
        """Contrôles *statiques* prudents, jamais une exécution du code.

        Le service ne connaît ni le projet ni la version cible de tous les
        langages. Il vérifie seulement ce qu'il peut réellement prouver.
        """
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ArtifactError(
                f"Le fichier source '{filename}' n'est pas en UTF-8."
            ) from error

        # Un bloc Markdown complet est un emballage, pas un fichier source.
        if text.lstrip().startswith("```"):
            raise ArtifactError(
                f"Le fichier source '{filename}' contient encore une "
                "délimitation Markdown au lieu du code brut."
            )

        if extension == "py" and text.strip():
            try:
                ast.parse(text, filename=filename)
            except SyntaxError as error:
                raise ArtifactError(
                    f"Syntaxe Python invalide dans '{filename}', "
                    f"ligne {error.lineno or '?'} : {error.msg}. "
                    "Aucune exécution n'a été effectuée."
                ) from error

    def _create_zip(
        self,
        content: Any,
        *,
        depth: int,
        strict_project: bool = False,
    ) -> bytes:
        if depth >= self.MAX_ARCHIVE_DEPTH:
            raise ArtifactError(
                "Profondeur maximale des archives ZIP atteinte."
            )

        if not isinstance(content, dict):
            raise ArtifactError(
                "Une archive ZIP attend un objet {'files': [...]}."
            )

        files = content.get("files")

        if not isinstance(files, list):
            raise ArtifactError(
                "Une archive ZIP attend une liste 'files'."
            )

        if strict_project and not files:
            raise ArtifactError(
                "Un ZIP de projet ORIA ne peut pas être vide."
            )

        if len(files) > self.MAX_ARCHIVE_FILES:
            raise ArtifactError(
                f"Une archive Oria accepte au maximum "
                f"{self.MAX_ARCHIVE_FILES} fichiers."
            )

        # Un éventuel inventaire explicite permet au futur orchestrateur de
        # refuser une livraison partielle. Sans inventaire, aucune promesse
        # d'exhaustivité des fichiers demandés n'est possible.
        expected_names: set[str] | None = None
        if strict_project and "expected_files" in content:
            expected = content["expected_files"]
            if (not isinstance(expected, list) or not expected
                    or len(expected) > self.MAX_ARCHIVE_FILES):
                raise ArtifactError("Inventaire expected_files invalide.")
            expected_names = set()
            expected_casefold: set[str] = set()
            for name in expected:
                if not isinstance(name, str):
                    raise ArtifactError("Un nom attendu doit être du texte.")
                safe = self._safe_archive_name(name)
                self._validate_project_archive_name(name, safe)
                if safe.casefold() in expected_casefold:
                    raise ArtifactError(
                        f"Nom dupliqué dans l'inventaire : {name!r}."
                    )
                expected_names.add(safe)
                expected_casefold.add(safe.casefold())

        buffer = io.BytesIO()

        with zipfile.ZipFile(
            buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=6,
        ) as archive:
            used_names: set[str] = set()
            used_code_names_casefold: set[str] = set()
            used_project_names_casefold: set[str] = set()
            total_uncompressed_bytes = 0

            for index, item in enumerate(files, start=1):
                if not isinstance(item, dict):
                    raise ArtifactError(
                        f"Entrée ZIP #{index} invalide."
                    )

                item_name = str(
                    item.get("name")
                    or item.get("filename")
                    or f"fichier-{index}.txt"
                )
                archive_name = self._safe_archive_name(item_name)
                item_type = str(
                    item.get("file_type")
                    or item.get("type")
                    or PurePosixPath(archive_name).suffix.lstrip(".")
                    or "txt"
                )
                normalized_type = self._normalize_type(item_type, archive_name)
                is_code_source = normalized_type in self.CODE_SOURCE_TYPES

                # Dans un projet issu du Chat, TOUT fichier compte :
                # package.json, README.md, .env, Dockerfile, etc. Un renommage
                # automatique peut briser l'exécution autant qu'un .py.
                if strict_project:
                    self._validate_project_archive_name(item_name, archive_name)
                    suffix = PurePosixPath(archive_name).suffix.lstrip(".").lower()
                    if suffix and suffix != normalized_type:
                        raise ArtifactError(
                            f"Extension incohérente pour '{archive_name}' : "
                            f"le contenu est déclaré '{normalized_type}'."
                        )
                    if archive_name.casefold() in used_project_names_casefold:
                        raise ArtifactError(
                            f"Fichier dupliqué dans le projet ZIP : {archive_name}."
                        )
                elif is_code_source:
                    # Compatibilité conservée pour les ZIP génériques.
                    self._validate_project_archive_name(item_name, archive_name)
                    declared_suffix = PurePosixPath(archive_name).suffix.lstrip(".")
                    if declared_suffix and declared_suffix != normalized_type:
                        raise ArtifactError(
                            f"Extension incohérente pour '{archive_name}' : "
                            f"le contenu est déclaré '{normalized_type}'."
                        )
                    casefold_name = archive_name.casefold()
                    if (
                        archive_name in used_names
                        or casefold_name in used_code_names_casefold
                    ):
                        raise ArtifactError(
                            f"Fichier source dupliqué dans le ZIP : "
                            f"{archive_name}."
                        )
                elif archive_name in used_names:
                    # Comportement historique conservé pour les autres ZIP.
                    archive_name = self._deduplicate_filename(
                        archive_name, used_names,
                    )

                item_content = item.get("content", "")
                if is_code_source and not isinstance(item_content, (str, bytes)):
                    raise ArtifactError(
                        f"Le fichier source '{archive_name}' attend du "
                        "texte ou des octets."
                    )

                if (is_code_source and item_content in ("", b"")
                        and PurePosixPath(archive_name).name == "__init__.py"):
                    # Seul __init__.py est volontairement accepté vide.
                    # Un main.py vide n'est pas un projet livré correctement.
                    generated_data = b""
                else:
                    item_metadata = item.get("metadata")
                    if strict_project and normalized_type == "zip":
                        item_metadata = {
                            **(item_metadata if isinstance(item_metadata, dict) else {}),
                            "source": "oria-chat",
                        }
                    generated = self.create(
                        file_type=item_type,
                        content=item_content,
                        filename=archive_name,
                        title=item.get("title"),
                        metadata=item_metadata,
                        allow_text_fallback=True,
                        _archive_depth=depth + 1,
                    )
                    generated_data = generated.data

                if strict_project:
                    if is_code_source and generated_data:
                        self._validate_chat_source(
                            filename=archive_name,
                            extension=normalized_type,
                            data=generated_data,
                        )
                    # Ne jamais transformer silencieusement un package.json
                    # invalide en chaîne JSON via _create_json().
                    if normalized_type == "json":
                        def forbid_non_json_constant(value: str) -> None:
                            raise ValueError(f"Constante JSON interdite : {value}")

                        # Le validateur porte sur les OCTETS livrés, pas
                        # seulement sur le texte initial : pas de NaN/Infinity.
                        try:
                            if isinstance(item_content, str):
                                json.loads(item_content, parse_constant=forbid_non_json_constant)
                            json.loads(
                                generated_data.decode("utf-8-sig"),
                                parse_constant=forbid_non_json_constant,
                            )
                        except json.JSONDecodeError as error:
                            raise ArtifactError(
                                f"JSON invalide dans '{archive_name}' : "
                                f"ligne {error.lineno}."
                            ) from error
                        except (ValueError, UnicodeDecodeError) as error:
                            raise ArtifactError(
                                f"JSON non conforme dans '{archive_name}' : {error}."
                            ) from error

                total_uncompressed_bytes += len(generated_data)
                if total_uncompressed_bytes > self.MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise ArtifactError(
                        "Le contenu décompressé de l'archive dépasse la "
                        "limite autorisée. Diviser le projet en plusieurs ZIP."
                    )

                archive.writestr(archive_name, generated_data)
                used_names.add(archive_name)
                if is_code_source:
                    used_code_names_casefold.add(archive_name.casefold())
                if strict_project:
                    used_project_names_casefold.add(archive_name.casefold())

        if strict_project:
            # 'src' ne peut pas être à la fois un fichier et le dossier qui
            # contient 'src/main.py' : certains extracteurs écraseraient l'un.
            for name in used_names:
                path_parts = name.split("/")
                for count in range(1, len(path_parts)):
                    if "/".join(path_parts[:count]).casefold() in used_project_names_casefold:
                        raise ArtifactError(
                            f"Conflit fichier/dossier dans le ZIP : '{name}'."
                        )

        if expected_names is not None:
            missing = sorted(expected_names - used_names)
            if missing:
                raise ArtifactError(
                    "Projet incomplet : fichiers absents du ZIP : "
                    + ", ".join(missing[:12])
                    + ("…" if len(missing) > 12 else "")
                )

        # Vérifie la table centrale et les noms effectivement archivés,
        # sans décompresser intégralement le projet (coût Render maîtrisé).
        if strict_project:
            try:
                with zipfile.ZipFile(io.BytesIO(buffer.getvalue()), "r") as check:
                    if set(check.namelist()) != used_names:
                        raise ArtifactError(
                            "Le ZIP final ne contient pas tous les fichiers annoncés."
                        )
            except (zipfile.BadZipFile, OSError) as error:
                raise ArtifactError("Le ZIP créé est illisible.") from error

        return buffer.getvalue()

    # ============================================================
    # ORIA ARTIFACT V1 — COMPOSITION DECLARATIVE SECURISEE
    # ============================================================
    # Chemin supplementaire et opt-in. create() reste retrocompatible.
    # Le protocole est aussi controle en amont par OpenAIService, mais
    # les donnees du modele demeurent non fiables ici.

    MAX_CONTRACT_SOURCE_CHARS = 1_000_000
    MAX_CONTRACT_ELEMENTS = 6000
    MAX_CONTRACT_ARTIFACTS = 15
    MAX_CONTRACT_IMAGE_BYTES = 2 * 1024 * 1024

    def create_from_payload(
        self, payload: dict[str, Any], *, metadata: dict[str, Any] | None = None
    ) -> list[ArtifactResult]:
        """Convertit une enveloppe oria.artifact.v1 en fichiers natifs.

        Tout est construit avant de retourner les resultats : aucun fichier
        n'est enregistre ici, le routeur decide de la persistance et de sa
        transaction. Aucun code de source issu du modele n'est execute.
        """
        if not isinstance(payload, dict) or set(payload) != {
            "schema_version", "sources", "artifacts"
        } or payload.get("schema_version") != "oria.artifact.v1":
            raise ArtifactError("Contrat documentaire oria.artifact.v1 invalide.")
        sources = payload["sources"]
        specs = payload["artifacts"]
        if not isinstance(sources, list) or len(sources) > 10:
            raise ArtifactError("Liste de sources documentaires invalide.")
        if not isinstance(specs, list) or not 1 <= len(specs) <= self.MAX_CONTRACT_ARTIFACTS:
            raise ArtifactError("Liste des fichiers a generer invalide.")
        permitted = {"pdf", "docx", "rtf", "odt", "epub", "pptx", "odp",
                     "xlsx", "ods", "csv", "tsv", "json", "jsonl", "ndjson", "sqlite"}
        safe_sources: dict[str, tuple[str, str]] = {}
        for source in sources:
            if not isinstance(source, dict) or set(source) != {"id", "kind", "content"}:
                raise ArtifactError("Source documentaire mal definie.")
            sid, kind, content = source["id"], source["kind"], source["content"]
            if (not isinstance(sid, str) or re.fullmatch(r"[A-Za-z0-9_-]{1,80}", sid) is None
                    or sid in safe_sources):
                raise ArtifactError("Identifiant de source invalide ou duplique.")
            if (kind not in {"html", "svg"} or not isinstance(content, str)
                    or not content.strip() or len(content) > self.MAX_CONTRACT_SOURCE_CHARS):
                raise ArtifactError("Source documentaire vide, trop grande ou non supportee.")
            if kind == "html":
                # Le modèle peut inclure des diagrammes SVG dans le HTML.
                # Les normaliser en images locales et vérifier chaque SVG
                # avant toute opération CairoSVG/WeasyPrint. Aucun chargement
                # distant, code actif ou accès au disque n'est autorisé.
                content = self._normalize_inline_svg_document(content)
                self._check_contract_html(content)
            else:
                self._check_contract_svg(content)
            safe_sources[sid] = (kind, content)

        generated: list[ArtifactResult] = []
        used_formats: set[str] = set()
        for spec in specs:
            if not isinstance(spec, dict) or set(spec) != {
                "format", "filename", "source_id", "native_data"
            }:
                raise ArtifactError("Specification de fichier documentaire invalide.")
            fmt, name = spec["format"], spec["filename"]
            sid, native = spec["source_id"], spec["native_data"]
            if not isinstance(fmt, str) or fmt not in permitted or fmt in used_formats:
                raise ArtifactError("Format documentaire non supporte ou duplique.")
            used_formats.add(fmt)
            if (not isinstance(name, str) or not 1 <= len(name) <= self.MAX_FILENAME_LENGTH
                    or not name.lower().endswith("." + fmt)
                    or name.startswith(".") or any(ch in name for ch in "/\\")
                    or any(ord(ch) < 32 for ch in name)
                    or self._resolve_filename(filename=name, title=None, extension=fmt) != name):
                raise ArtifactError("Nom documentaire dangereux ou modifie silencieusement.")
            if sid is not None and (not isinstance(sid, str) or sid not in safe_sources):
                raise ArtifactError("Reference de source documentaire invalide.")
            if native is not None and not isinstance(native, (dict, list)):
                raise ArtifactError("Les donnees natives doivent etre un objet ou tableau JSON.")
            if sid is None and native is None:
                raise ArtifactError("Fichier sans source ni donnees natives.")
            kind, source = safe_sources[sid] if sid is not None else (None, None)
            try:
                data = self._build_contract_bytes(fmt, kind, source, native, name)
            except ArtifactError:
                raise
            except (TypeError, ValueError, KeyError, OverflowError) as error:
                raise ArtifactError(f"Contenu invalide pour {fmt}: {error}") from error
            if not isinstance(data, bytes):
                raise ArtifactError("Le moteur documentaire n'a pas retourne des octets.")
            self._validate_size(data)
            result_meta = dict(metadata or {})
            result_meta.update({"format": fmt, "generator": "oria-artifact-v1",
                                "schema_version": "oria.artifact.v1"})
            generated.append(ArtifactResult(
                artifact_id=str(uuid4()), filename=name, extension=fmt,
                mime_type=self._resolve_mime_type(fmt),
                category=self.CATEGORY_BY_TYPE.get(fmt, "document"),
                data=data, size_bytes=len(data),
                sha256=hashlib.sha256(data).hexdigest(), metadata=result_meta,
            ))
        return generated

    @classmethod
    def _normalize_inline_svg_document(cls, source: str) -> str:
        """Accepte les diagrammes SVG inline *sûrs* dans les documents HTML.

        La représentation du modèle n'est jamais exécutée : chaque élément
        SVG est extrait sans modifier son XML (casse de viewBox comprise),
        validé avec le parseur XML sécurisé, puis rasterisé localement en PNG
        autonome. Les fichiers Word et PDF reçoivent ainsi le même diagramme.
        Le SVG non conforme est rejeté : jamais de suppression silencieuse.
        """
        if not re.search(r"<\s*svg\b", source, re.I):
            return source
        try:
            from bs4 import BeautifulSoup
        except ImportError as error:
            raise ArtifactDependencyError(
                "Illustrations SVG HTML : installer beautifulsoup4."
            ) from error
        soup = BeautifulSoup(source, "html.parser")
        svg_nodes = soup.find_all("svg")
        # Aucune prise en charge des SVG imbriqués. La substitution textuelle
        # doit retrouver exactement le nombre d'éléments du DOM, sinon rejet.
        if not 1 <= len(svg_nodes) <= 8:
            raise ArtifactError("Nombre d'illustrations SVG HTML non autorisé.")
        if any(node.find("svg") is not None for node in svg_nodes):
            raise ArtifactError("SVG imbriqués non pris en charge.")

        regex = re.compile(r"<svg\b[^>]*>[\s\S]*?</svg\s*>", re.I)
        if len(regex.findall(source)) != len(svg_nodes):
            raise ArtifactError("Illustration SVG HTML incomplète ou mal formée.")
        try:
            import cairosvg
            from lxml import etree
        except (ImportError, OSError) as error:
            raise ArtifactDependencyError(
                "Illustrations SVG HTML : CairoSVG et lxml sont requis."
            ) from error
        import base64
        from math import sqrt

        def replace(match: re.Match[str]) -> str:
            svg = match.group(0)
            cls._check_contract_svg(svg)
            root = etree.fromstring(svg.encode("utf-8"), parser=etree.XMLParser(
                resolve_entities=False, load_dtd=False, no_network=True,
                huge_tree=False,
            ))
            box = root.get("viewBox")
            if box:
                values = [float(x) for x in re.split(r"[\s,]+", box.strip())]
                width, height = values[2], values[3]
            else:
                def numeric_size(name: str, default: float) -> float:
                    text = root.get(name) or ""
                    m = re.fullmatch(r"(\d+(?:\.\d+)?)(?:px|pt|mm|cm)?", text.strip())
                    return float(m.group(1)) if m else default
                width, height = numeric_size("width", 400), numeric_size("height", 240)
            if width <= 0 or height <= 0 or max(width / height, height / width) > 40:
                raise ArtifactError("Proportions SVG HTML non autorisées.")
            scale = min(1400 / width, 1400 / height,
                        sqrt(1_500_000 / (width * height)))
            out_width = max(1, round(width * scale))
            out_height = max(1, round(height * scale))
            try:
                png = cairosvg.svg2png(
                    bytestring=svg.encode("utf-8"),
                    output_width=out_width, output_height=out_height,
                )
            except Exception as error:
                raise ArtifactError(
                    "Conversion de l'illustration SVG intégrée impossible."
                ) from error
            if (not isinstance(png, bytes) or not png.startswith(b"\x89PNG\r\n\x1a\n")
                    or len(png) > cls.MAX_CONTRACT_IMAGE_BYTES):
                raise ArtifactError("Illustration SVG convertie trop volumineuse.")
            title = next((" ".join(node.itertext()).strip() for node in root
                          if etree.QName(node).localname in {"title", "desc"}), "")
            # Une légende alternative courte seulement, aucun attribut SVG
            # ni contenu XML non fiable n'est injecté dans le DOM HTML final.
            safe_alt = html_escape(title[:150] or "Illustration vectorielle", quote=True)
            data = base64.b64encode(png).decode("ascii")
            # Fixer la largeur d'affichage : sans contrainte CSS, WeasyPrint
            # place le bitmap a sa taille intrinseque et peut le couper au
            # bord de la page A4.
            return (
                f'<img alt="{safe_alt}" '
                'style="display:block;max-width:100%;height:auto" '
                f'src="data:image/png;base64,{data}">'
            )

        normalized, count = regex.subn(replace, source)
        if count != len(svg_nodes) or len(normalized) > 9_000_000:
            raise ArtifactError("Document HTML avec illustrations SVG trop volumineux.")
        return normalized

    @classmethod
    def _check_contract_html(cls, source: str) -> None:
        """Interdit HTML actif et ressources externes (double barriere fetcher)."""
        try:
            from bs4 import BeautifulSoup
        except ImportError as error:
            raise ArtifactDependencyError("Composition HTML : installer beautifulsoup4.") from error
        if not re.search(r"<(?:html|body|main)\b", source, flags=re.I):
            raise ArtifactError("La source HTML ne comporte pas de document valide.")
        if re.search(r"<!\s*(?:DOCTYPE\s+[^>]*\[|ENTITY)\b", source, flags=re.I):
            raise ArtifactError("Declarations de document dangereuses interdites.")
        soup = BeautifulSoup(source, "html.parser")
        tags = soup.find_all(True)
        if len(tags) > cls.MAX_CONTRACT_ELEMENTS:
            raise ArtifactError("Document HTML trop complexe pour le moteur.")
        allowed = {
            "html", "head", "body", "meta", "title", "style", "main", "section",
            "article", "div", "header", "footer", "h1", "h2", "h3", "h4", "h5", "h6",
            "p", "span", "strong", "em", "b", "i", "u", "s", "small", "sup", "sub",
            "blockquote", "ul", "ol", "li", "table", "thead", "tbody", "tfoot",
            "tr", "td", "th", "caption", "br", "hr", "pre", "code", "img", "figure",
            "figcaption", "colgroup", "col", "a"
        }
        allowed_attrs = {"class", "id", "style", "lang", "dir", "title", "alt",
                         "src", "width", "height", "colspan", "rowspan", "scope",
                         "data-type", "charset", "name", "content", "href",
                         "role", "aria-label", "aria-hidden", "aria-describedby",
                         "aria-labelledby"}
        for tag in tags:
            if tag.name not in allowed:
                raise ArtifactError(f"Balise HTML documentaire interdite : {tag.name}.")
            for attribute, value in tag.attrs.items():
                if attribute not in allowed_attrs:
                    raise ArtifactError(f"Attribut HTML non autorise : {attribute}.")
                if attribute == "href":
                    if tag.name != "a" or not isinstance(value, str) or not (value.startswith("#") or re.fullmatch(r"https://[^\s<>\"\']{1,1000}", value, re.I)):
                        raise ArtifactError("Lien HTML non autorise.")
                elif attribute == "src":
                    if tag.name != "img" or not isinstance(value, str):
                        raise ArtifactError("Ressource externe interdite.")
                    cls._image_data_uri(value)
                elif attribute == "style":
                    cls._check_contract_css(str(value))
                elif attribute in {"width", "height", "colspan", "rowspan"}:
                    if not str(value).isdigit() or int(value) > 3000:
                        raise ArtifactError("Dimension HTML non autorisee.")
            if tag.name == "meta" and tag.get("http-equiv"):
                raise ArtifactError("Redirection HTML interdite.")
            if tag.name == "style":
                cls._check_contract_css(tag.get_text())
        # Les seuls attributs meta effectivement utiles restent fixes.
        for meta in soup.find_all("meta"):
            if set(meta.attrs) - {"charset", "name", "content"}:
                raise ArtifactError("Meta HTML non autorise.")
            if (meta.get("name") or "").lower() not in {"", "viewport", "description"}:
                raise ArtifactError("Meta HTML non autorise.")

    @staticmethod
    def _check_contract_css(css: str) -> None:
        # Interdire tout chargement distant/local via CSS, independamment
        # du moteur de rendu. Aucun Javascript ni CSS dynamique.
        if len(css) > 180_000 or re.search(
            r"@import\b|@font-face\b|url\s*\(|expression\s*\(|behavior\s*:|"
            r"-moz-binding|javascript\s*:|data\s*:", css, re.I
        ):
            raise ArtifactError("CSS externe, actif ou trop volumineux interdit.")

    @classmethod
    def _image_data_uri(cls, raw: str) -> tuple[str, bytes]:
        match = re.fullmatch(r"data:image/(png|jpeg|webp|gif);base64,([A-Za-z0-9+/=\r\n]+)", raw, re.I)
        if not match or len(raw) > cls.MAX_CONTRACT_IMAGE_BYTES * 2:
            raise ArtifactError("Seules les images data: PNG/JPEG/WebP/GIF limitees sont admises.")
        import base64
        try:
            data = base64.b64decode(match.group(2), validate=True)
        except (ValueError, base64.binascii.Error) as error:
            raise ArtifactError("Image integree mal encodee.") from error
        if not data or len(data) > cls.MAX_CONTRACT_IMAGE_BYTES:
            raise ArtifactError("Image integree trop volumineuse.")
        # Verifier les dimensions avant que les moteurs n'allouent un bitmap.
        try:
            from PIL import Image
            with Image.open(io.BytesIO(data)) as image:
                if image.width * image.height > 8_000_000:
                    raise ArtifactError("Image integree trop grande.")
                image.verify()
        except ImportError:
            pass
        except ArtifactError:
            raise
        except Exception as error:
            raise ArtifactError("Image integree illisible.") from error
        return match.group(1).lower(), data

    @classmethod
    def _check_contract_svg(cls, source: str) -> None:
        if "<!DOCTYPE" in source.upper() or "<!ENTITY" in source.upper() or len(source) > cls.MAX_CONTRACT_SOURCE_CHARS:
            raise ArtifactError("Declaration SVG externe ou dangereuse interdite.")
        try:
            from lxml import etree
        except ImportError as error:
            raise ArtifactDependencyError("SVG : installer lxml.") from error
        try:
            root = etree.fromstring(source.encode("utf-8"), parser=etree.XMLParser(
                resolve_entities=False, load_dtd=False, no_network=True, huge_tree=False
            ))
        except etree.XMLSyntaxError as error:
            raise ArtifactError("Source SVG non conforme.") from error
        if etree.QName(root).localname != "svg":
            raise ArtifactError("La source SVG doit avoir une racine svg.")
        for attr in ('width','height'):
            value=root.get(attr)
            if value is not None:
                mm=re.fullmatch(r'([0-9]{1,5}(?:\.[0-9]{1,3})?)(?:px|pt|mm|cm)?',value.strip())
                if mm is None or float(mm.group(1))>5000:
                    raise ArtifactError('Dimensions SVG non autorisees.')
        view_box=root.get('viewBox')
        if view_box is not None:
            try:
                dims=[float(i) for i in re.split(r'[\s,]+', view_box.strip())]
                if len(dims)!=4 or not 0<dims[2]<=5000 or not 0<dims[3]<=5000 or max(dims[2]/dims[3],dims[3]/dims[2])>40:
                    raise ValueError('dimensions incoherentes')
            except (ValueError,OverflowError,ZeroDivisionError) as error:
                raise ArtifactError('viewBox SVG non autorise.') from error
        allowed = {"svg", "g", "defs", "path", "rect", "circle", "ellipse", "line",
                   "polyline", "polygon", "text", "tspan", "title", "desc", "linearGradient",
                   "radialGradient", "stop", "clipPath", "use"}
        for index, el in enumerate(root.iter()):
            if index > cls.MAX_CONTRACT_ELEMENTS:
                raise ArtifactError("SVG trop volumineux.")
            if not isinstance(el.tag, str) or etree.QName(el).localname not in allowed:
                raise ArtifactError("Element SVG non autorise.")
            for attr, val in el.attrib.items():
                a = etree.QName(attr).localname.lower()
                if a.startswith("on") or a in {"href", "src", "filter"}:
                    raise ArtifactError("Attribut SVG actif ou externe interdit.")
                if "url(" in val.lower() and not re.fullmatch(r"url\(\s*#[A-Za-z0-9_-]+\s*\)", val, re.I):
                    raise ArtifactError("Reference externe dans SVG interdite.")
                if "javascript:" in val.lower() or "@import" in val.lower():
                    raise ArtifactError("Valeur SVG dangereuse.")

    @staticmethod
    def _blocked_document_resource(url: str, *args: Any, **kwargs: Any) -> dict[str, Any]:
        # WeasyPrint invoque ce callback pour CHAQUE ressource incluse. Aucune
        # lecture du filesystem, du reseau ou d'URL file:// n'est permise.
        if not url.lower().startswith("data:"):
            raise ArtifactError("Chargement externe interdit pendant le rendu.")
        mime, data = ArtifactService._image_data_uri(url)
        return {"string": data, "mime_type": "image/" + mime, "redirected_url": url}

    def _build_contract_bytes(
        self, fmt: str, kind: str | None, source: str | None,
        native: dict[str, Any] | list[Any] | None, filename: str
    ) -> bytes:
        # Ne jamais inserer silencieusement un objet JSON comme du texte
        # dans un document : seuls les plans de composition reconnus sont admis.
        if fmt in {'pdf','docx','rtf','odt','epub'} and source is None:
            if not isinstance(native,dict) or not ({'blocks','chapters'} & set(native)):
                raise ArtifactError(
                    f'{fmt.upper()} attend native_data.blocks ou native_data.chapters.')
        if fmt == 'odp' and isinstance(native,dict) and isinstance(native.get('slides'),list):
            if any(set(slide) & {'table','chart','elements','subtitle'}
                   for slide in native['slides'] if isinstance(slide,dict)):
                raise UnsupportedArtifactTypeError(
                    'Composition ODP avancee non encore disponible : eviter les pertes silencieuses.')
        if kind == 'svg' and fmt not in {'pdf', 'docx'} and native is None:
            raise UnsupportedArtifactTypeError(
                f'Une source SVG ne peut pas encore etre convertie fidelement en {fmt}.')
        if (kind == 'html' and source and re.search(r'<img\b', source, re.I)
                and fmt in {'odt', 'rtf', 'pptx', 'odp'} and native is None):
            raise UnsupportedArtifactTypeError(
                f'Les illustrations HTML ne sont pas encore conservees en {fmt}.')
        if fmt == "pdf":
            return self._contract_pdf(kind, source, native, filename)
        if fmt == "docx":
            return self._contract_docx(kind, source, native)
        if fmt == "odt":
            return self._contract_odt(kind, source, native)
        if fmt == "rtf":
            return self._contract_rtf(kind, source, native)
        if fmt == "epub":
            return self._contract_epub(kind, source, native, filename)
        if fmt == "pptx":
            return self._contract_pptx(kind, source, native)
        if fmt == "odp":
            slides = self._contract_slides(kind, source, native)
            return self._create_odp({"slides": slides}, title=None)
        if fmt == "xlsx":
            return self._contract_xlsx(kind, source, native)
        if fmt == "ods":
            return self._contract_ods(kind, source, native)
        if fmt in {"csv", "tsv"}:
            return self._contract_delimited(fmt, kind, source, native)
        if fmt in {"json", "jsonl", "ndjson"}:
            if native is None:
                raise ArtifactError(f"{fmt.upper()} necessite des donnees natives, pas du HTML.")
            self._require_json_serializable(native)
            if fmt == "json":
                return self._create_json(native)
            return self._create_jsonl(native)
        if fmt == "sqlite":
            if native is None:
                raise ArtifactError("SQLite necessite des tables et donnees natives.")
            return self._contract_sqlite(native)
        raise UnsupportedArtifactTypeError(f"Conversion indisponible : {fmt}.")

    @staticmethod
    def _require_json_serializable(value: Any) -> None:
        try:
            json.dumps(value, allow_nan=False)
        except (TypeError, ValueError) as error:
            raise ArtifactError("Donnees natives non serialisables en JSON.") from error

    def _contract_document_html(self, native: Any) -> str:
        blocks = self._contract_blocks(None, None, native)
        markup: list[str] = []
        for block in blocks:
            kind = block.get("type", "p")
            label = html_escape(str(block.get("text", "")))
            if kind in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                markup.append(f"<{kind}>{label}</{kind}>")
            elif kind == "page_break":
                markup.append('<div class="page-break"></div>')
            elif kind == "bullet":
                items = "".join(f"<li>{html_escape(str(v))}</li>" for v in block.get("items", []))
                markup.append("<ul>" + items + "</ul>")
            elif kind == "table":
                markup.append("<table>" + "".join("<tr>" + "".join(
                    f"<td>{html_escape(str(v if v is not None else ''))}</td>" for v in row
                ) + "</tr>" for row in block.get("rows", [])) + "</table>")
            else:
                markup.append(f"<p>{label}</p>")
        body = "\n".join(markup)
        return ('<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8">'
                '<style>@page{size:A4;margin:20mm}body{font:11pt/1.55 Arial,sans-serif;color:#202d40}'
                'h1,h2,h3{color:#17356d;break-after:avoid}p{orphans:3;widows:3}'
                '.page-break{break-before:page}table{border-collapse:collapse;width:100%}'
                'td,th{border:1px solid #ccd5e0;padding:6px}</style></head><body><main>'
                + body + '</main></body></html>')

    def _contract_pdf(self, kind: str | None, source: str | None,
                      native: Any, filename: str) -> bytes:
        if source is None:
            if native is None:
                raise ArtifactError("PDF sans contenu.")
            source = self._contract_document_html(native)
            kind = "html"
        if kind == "svg":
            try:
                import cairosvg
            except ImportError as error:
                raise ArtifactDependencyError('SVG vectoriel vers PDF : installer cairosvg.') from error
            # CairoSVG conserve ici les traces vectorielles dans le PDF.
            # La source a ete analysee par _check_contract_svg auparavant.
            try:
                return cairosvg.svg2pdf(bytestring=source.encode('utf-8'))
            except Exception as error:
                raise ArtifactError('Echec du rendu SVG vectoriel vers PDF.') from error
        self._check_contract_html(source)
        try:
            from weasyprint import HTML
        except (ImportError, OSError) as error:
            # Une installation pip peut réussir alors que Pango/GLib/Cairo
            # manque encore dans l'image système de Render/Windows.
            raise ArtifactDependencyError(
                "Rendu PDF HTML/CSS indisponible : verifier WeasyPrint, Pango et GLib."
            ) from error
        try:
            return HTML(string=source, base_url=None,
                        url_fetcher=self._blocked_document_resource).write_pdf()
        except ArtifactError:
            raise
        except OSError as error:
            raise ArtifactDependencyError(
                "Le rendu PDF ne trouve pas une bibliotheque systeme de typographie."
            ) from error
        except Exception as error:
            # L'exception d'origine demeure attachée au traceback, visible
            # via ORIA_ARTIFACT_FAILURE/ref dans les logs, jamais dans le Chat.
            raise ArtifactError("Echec du rendu HTML/CSS vers PDF.") from error

    @classmethod
    def _contract_html_blocks(cls, html: str) -> list[dict[str, Any]]:
        from bs4 import BeautifulSoup, Tag
        soup = BeautifulSoup(html, "html.parser")
        root = soup.find("main") or soup.body or soup
        blocks: list[dict[str, Any]] = []
        def visit(node: Any) -> None:
            if not isinstance(node, Tag):
                return
            name = node.name
            classes = node.get('class') or []
            raw_style = str(node.get('style') or '').lower()
            if ("page-break" in classes or 'break-before:page' in raw_style.replace(' ', '')
                    or 'page-break-before:always' in raw_style.replace(' ', '')):
                blocks.append({"type":"page_break"})
            if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                blocks.append({"type": name, "text": node.get_text(" ", strip=True), "element": node})
            elif name in {"p", "blockquote", "pre", "figcaption"}:
                content = node.get_text(" ", strip=True)
                if content:
                    blocks.append({"type":"p", "text":content, "element":node})
            elif name in {"ul", "ol"}:
                items = [li for li in node.find_all('li', recursive=False)]
                if items:
                    blocks.append({"type":"number" if name == "ol" else "bullet",
                                   "items":[li.get_text(" ", strip=True) for li in items], "elements":items})
            elif name == "table":
                rows = [[cell.get_text(" ", strip=True) for cell in tr.find_all(["th","td"], recursive=False)]
                        for tr in node.find_all("tr")]
                if rows:
                    blocks.append({"type":"table", "rows":rows})
            elif name == "img":
                blocks.append({"type":"image", "src": node.get("src"), "alt":node.get("alt", "")})
            elif name == "br":
                return
            else:
                children = [child for child in node.children if isinstance(child, Tag)]
                if children and all(c.name in {'strong','em','b','i','u','span','small','sup','sub','a','code'} for c in children):
                    content=node.get_text(' ',strip=True)
                    if content:blocks.append({'type':'p','text':content,'element':node})
                    return
                direct_text = " ".join(ch.strip() for ch in node.children if isinstance(ch, str) and ch.strip())
                if direct_text:
                    blocks.append({"type":"p", "text":direct_text})
                for child in children:
                    visit(child)
        for child in root.children:
            visit(child)
        return blocks

    def _contract_blocks(self, kind: str | None, source: str | None, native: Any) -> list[dict[str, Any]]:
        if native is not None and (source is None or kind == "svg"):
            if isinstance(native,dict) and not ({'blocks','chapters'} & set(native)):
                raise ArtifactError('Structure documentaire native inconnue.')
            if isinstance(native, dict) and isinstance(native.get('chapters'),list):
                blocks=[]
                for chapter in native['chapters']:
                    if not isinstance(chapter, dict):
                        raise ArtifactError('Chapitre documentaire invalide.')
                    blocks.append({'type':'h1','text':str(chapter.get('title') or 'Chapitre')})
                    content=chapter.get('blocks',chapter.get('content',chapter.get('text','')))
                    if isinstance(content,list):content={'blocks':content}
                    blocks.extend(self._normalize_document_blocks(content))
                return blocks
            blocks=self._normalize_document_blocks(native)
            supported={'p','h1','h2','h3','h4','h5','h6','bullet','number','table','page_break','spacer'}
            for b in blocks:
                if b['type'] not in supported:
                    raise ArtifactError('Bloc documentaire natif non supporte : '+b['type'])
            return blocks
        if source is None:
            raise ArtifactError("Document sans source interpretable.")
        if kind == "svg":
            return [{"type": "svg", "source": source}]
        return self._contract_html_blocks(source)

    def _contract_docx(self, kind: str | None, source: str | None, native: Any) -> bytes:
        try:
            from docx import Document
            from docx.shared import Cm, Pt, RGBColor
        except ImportError as error:
            raise ArtifactDependencyError("DOCX : installer python-docx.") from error
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.25); section.right_margin = Cm(2.25)
        normal = document.styles['Normal']
        normal.font.name = 'Aptos'; normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.25
        for level in (1, 2, 3):
            style = document.styles[f'Heading {level}']
            style.font.name = 'Aptos Display'
            style.font.color.rgb = RGBColor(23, 53, 109)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        blocks = self._contract_blocks(kind, source, native)
        def rich_runs(paragraph: Any, tag: Any, bold: bool = False,
                      italic: bool = False, underline: bool = False) -> None:
            from bs4 import NavigableString, Tag
            for child in tag.children:
                if isinstance(child, NavigableString):
                    if str(child):
                        run = paragraph.add_run(str(child))
                        run.bold = bold; run.italic = italic; run.underline = underline
                elif isinstance(child, Tag):
                    if child.name == 'br':
                        paragraph.add_run().add_break()
                    else:
                        rich_runs(paragraph, child, bold or child.name in {'b','strong'},
                                  italic or child.name in {'i','em'},
                                  underline or child.name == 'u')
        for block in blocks:
            typ = block['type']
            if typ == 'page_break':
                document.add_page_break()
            elif typ.startswith('h') and typ[1:].isdigit():
                para = document.add_heading(level=min(int(typ[1:]), 4))
                if block.get('element'):
                    rich_runs(para, block['element'])
                else:
                    para.add_run(block.get('text', ''))
            elif typ == 'p':
                para = document.add_paragraph()
                if block.get('element'):
                    rich_runs(para, block['element'])
                else:
                    para.add_run(block.get('text', ''))
            elif typ in {'bullet', 'number'}:
                for index, item in enumerate(block.get('items', [])):
                    para = document.add_paragraph(style='List Number' if typ == 'number' else 'List Bullet')
                    elements = block.get('elements') or []
                    if index < len(elements):
                        rich_runs(para, elements[index])
                    else:
                        para.add_run(str(item))
            elif typ == 'table':
                rows = block.get('rows') or []
                if rows:
                    numcols = max(1, max(len(row) for row in rows))
                    table = document.add_table(rows=len(rows), cols=numcols)
                    table.style = 'Light Shading Accent 1'
                    for r, cells in enumerate(rows):
                        for c, value in enumerate(cells):
                            table.cell(r, c).text = '' if value is None else str(value)
                    for cell in table.rows[0].cells:
                        for para in cell.paragraphs:
                            for run in para.runs: run.bold = True
            elif typ == 'image':
                _, data = self._image_data_uri(block['src'])
                document.add_picture(io.BytesIO(data), width=Cm(14))
            elif typ == 'svg':
                try:
                    import cairosvg
                except ImportError as error:
                    raise ArtifactDependencyError('SVG vers Word : installer cairosvg.') from error
                png = cairosvg.svg2png(bytestring=block['source'].encode(), output_width=1100)
                document.add_picture(io.BytesIO(png), width=Cm(14))
        if not blocks:
            raise ArtifactError('Le document Word ne contient aucun contenu exploitable.')
        buffer = io.BytesIO(); document.save(buffer)
        return buffer.getvalue()

    def _contract_odt(self, kind: str | None, source: str | None, native: Any) -> bytes:
        try:
            from odf.opendocument import OpenDocumentText
            from odf.text import H, P
            from odf.table import Table, TableRow, TableCell
        except ImportError as error:
            raise ArtifactDependencyError('ODT : installer odfpy.') from error
        doc = OpenDocumentText()
        blocks = self._contract_blocks(kind, source, native)
        if any(b['type'] in {'image','svg'} for b in blocks):
            raise UnsupportedArtifactTypeError('ODT avec illustration necessite un moteur image supplementaire.')
        for block in blocks:
            typ = block['type']; value = block.get('text','')
            if typ.startswith('h') and typ[1:].isdigit():
                doc.text.addElement(H(outlinelevel=min(int(typ[1:]), 6),text=value))
            elif typ in {'bullet', 'number'}:
                prefix = '• ' if typ == 'bullet' else ''
                for idx, item in enumerate(block.get('items', []),start=1):
                    doc.text.addElement(P(text=prefix + (f'{idx}. ' if typ=='number' else '') + str(item)))
            elif typ == 'table':
                tbl = Table(name='Tableau')
                for cells in block.get('rows', []):
                    row = TableRow()
                    for value in cells:
                        cell = TableCell(valuetype='string')
                        cell.addElement(P(text='' if value is None else str(value)))
                        row.addElement(cell)
                    tbl.addElement(row)
                doc.text.addElement(tbl)
            elif typ == 'page_break':
                # ODF requiert une propriete de style dediee pour un vrai saut.
                from odf.style import Style, ParagraphProperties
                style_name = 'OriaPageBreak'
                if not any(x.getAttribute('name') == style_name for x in doc.automaticstyles.childNodes if x.qname[1]=='style'):
                    style = Style(name=style_name, family='paragraph')
                    style.addElement(ParagraphProperties(breakbefore='page'))
                    doc.automaticstyles.addElement(style)
                doc.text.addElement(P(stylename=style_name,text=''))
            elif typ == 'p':
                doc.text.addElement(P(text=value))
        if not blocks:
            raise ArtifactError('Document ODT vide.')
        buffer = io.BytesIO(); doc.write(buffer); return buffer.getvalue()

    def _contract_rtf(self, kind: str | None, source: str | None, native: Any) -> bytes:
        blocks = self._contract_blocks(kind, source, native)
        if any(b['type'] in {'image','svg'} for b in blocks):
            raise UnsupportedArtifactTypeError('RTF avec illustration non pris en charge.')
        def esc(s: Any) -> str:
            out=[]
            for ch in str(s):
                val=ord(ch)
                if ch in '{}\\': out.append('\\'+ch)
                elif ch=='\n':out.append('\\line ')
                elif val > 127:
                    for unit in ([(val - 0x10000) // 1024 + 0xD800,
                                  (val - 0x10000) % 1024 + 0xDC00] if val>65535 else [val]):
                        out.append(f'\\u{unit if unit<32768 else unit-65536}?')
                else:out.append(ch)
            return ''.join(out)
        parts=[r'{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}\fs22 ']
        for block in blocks:
            typ=block['type']
            if typ.startswith('h') and typ[1:].isdigit():
                pts = {1:34, 2:28, 3:24}.get(int(typ[1:]),22)
                parts.append(f'\\pard\\sb240\\sa120\\b\\fs{pts} {esc(block.get("text",""))}\\b0\\fs22\\par\n')
            elif typ=='page_break':parts.append('\\page\n')
            elif typ in {'bullet','number'}:
                for i,v in enumerate(block.get('items', []),1):
                    parts.append('\\pard\\li500 '+esc(('• ' if typ=='bullet' else f'{i}. ')+str(v))+'\\par\n')
            elif typ=='table':
                for row in block.get('rows', []):
                    parts.append('\\pard '+esc(' | '.join(str(v or '') for v in row))+'\\par\n')
            elif typ=='p':parts.append('\\pard '+esc(block.get('text',''))+'\\par\n')
        parts.append('}')
        return ''.join(parts).encode('ascii')

    def _contract_epub(self, kind: str | None, source: str | None,
                       native: Any, filename: str) -> bytes:
        # Chaque titre h1 ouvre un chapitre XHTML, qui reste lisible/editable.
        blocks = self._contract_blocks(kind, source, native)
        chapters: list[dict[str, Any]]=[]
        epub_images: list[tuple[str, str, bytes]] = []
        current={"title":"Introduction", "blocks":[]}
        for block in blocks:
            if block['type']=='h1':
                if current['blocks']:chapters.append(current)
                current={"title":block.get('text','Chapitre') or 'Chapitre',"blocks":[]}
            else:current['blocks'].append(block)
        if current['blocks'] or not chapters:chapters.append(current)
        if not any(c['blocks'] for c in chapters):
            raise ArtifactError('EPUB sans texte ni illustration.')
        from datetime import datetime, timezone
        book_id=str(uuid4())
        modified = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
        entries=[]; nav=[]; spine=[]
        for idx,chap in enumerate(chapters,1):
            inner=[]
            for block in chap['blocks']:
                kind=block['type']
                if kind in {'h2','h3','h4','h5','h6','p'}:
                    tag=kind if kind!='p' else 'p'
                    inner.append(f'<{tag}>{html_escape(block.get("text",""))}</{tag}>')
                elif kind in {'bullet','number'}:
                    tag='ol' if kind=='number' else 'ul'
                    inner.append(f'<{tag}>'+''.join('<li>'+html_escape(str(v))+'</li>' for v in block.get('items',[]))+f'</{tag}>')
                elif kind=='table':
                    inner.append('<table>'+''.join('<tr>'+''.join('<td>'+html_escape(str(v or ''))+'</td>' for v in row)+'</tr>' for row in block.get('rows',[]))+'</table>')
                elif kind=='image':
                    mime, img = self._image_data_uri(str(block['src']))
                    if mime == 'webp':
                        raise UnsupportedArtifactTypeError('Convertir les images WebP en PNG avant EPUB.')
                    image_id=f'img{len(epub_images)+1}'
                    suffix={'jpeg':'jpg', 'png':'png', 'gif':'gif'}[mime]
                    image_name=image_id+'.'+suffix
                    epub_images.append((image_name,'image/'+mime,img))
                    inner.append('<figure><img src="'+image_name+'" alt="'+
                                 html_escape(str(block.get('alt','')),quote=True) + '" /></figure>')
            chapter_title=html_escape(chap['title'])
            chapter_data=('<?xml version="1.0" encoding="utf-8"?>'
                '<html xmlns="http://www.w3.org/1999/xhtml" lang="fr"><head><title>'
                + chapter_title + '</title><link rel="stylesheet" type="text/css" href="style.css"/>'
                '</head><body><h1>'+chapter_title+'</h1>'+''.join(inner)+'</body></html>')
            path=f'chapter{idx}.xhtml'
            entries.append((path,chapter_data))
            nav.append(f'<li><a href="{path}">{chapter_title}</a></li>')
            spine.append(f'<itemref idref="chap{idx}"/>')
        title=html_escape(PurePosixPath(filename).stem.replace('_',' '))
        package=('<?xml version="1.0" encoding="utf-8"?>'
            '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            f'<dc:identifier id="book-id">{book_id}</dc:identifier><dc:title>{title}</dc:title>'
            '<dc:language>fr</dc:language>'
            f'<meta property="dcterms:modified">{modified}</meta>'
            '</metadata><manifest>'
            '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
            '<item id="style" href="style.css" media-type="text/css"/>'
            + ''.join(f'<item id="chap{i}" href="chapter{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(1,len(entries)+1))
            + ''.join(f'<item id="im{i}" href="{name}" media-type="{mime}"/>'
                      for i,(name,mime,_) in enumerate(epub_images,1))
            + '</manifest><spine>'+''.join(spine)+'</spine></package>')
        toc=('<?xml version="1.0" encoding="utf-8"?>'
             '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="fr">'
             '<head><title>Sommaire</title></head><body><nav epub:type="toc" id="toc">'
             '<h1>Sommaire</h1><ol>'+''.join(nav)+'</ol></nav></body></html>')
        container=('<?xml version="1.0" encoding="utf-8"?>'
                   '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="EPUB/content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>')
        css='body{font-family:serif;line-height:1.6;color:#222}h1,h2{color:#183767}p{text-align:justify}table{border-collapse:collapse}td{border:1px solid #aaa;padding:.3em}'
        output=io.BytesIO()
        with zipfile.ZipFile(output,'w') as archive:
            archive.writestr('mimetype','application/epub+zip',compress_type=zipfile.ZIP_STORED)
            archive.writestr('META-INF/container.xml',container,compress_type=zipfile.ZIP_DEFLATED)
            archive.writestr('EPUB/content.opf',package,compress_type=zipfile.ZIP_DEFLATED)
            archive.writestr('EPUB/nav.xhtml',toc,compress_type=zipfile.ZIP_DEFLATED)
            archive.writestr('EPUB/style.css',css,compress_type=zipfile.ZIP_DEFLATED)
            for fn,data in entries:
                archive.writestr('EPUB/'+fn,data,compress_type=zipfile.ZIP_DEFLATED)
            for fn,mime,data in epub_images:
                archive.writestr('EPUB/'+fn,data,compress_type=zipfile.ZIP_DEFLATED)
        return output.getvalue()

    def _contract_slides(self, kind: str | None, source: str | None,
                         native: Any) -> list[dict[str, Any]]:
        if isinstance(native, dict) and isinstance(native.get('slides'), list):
            slides = native['slides']
            if not slides or len(slides)>50 or any(not isinstance(slide, dict) for slide in slides):
                raise ArtifactError('Liste de diapositives invalide (1-50).')
            return slides
        blocks=self._contract_blocks(kind,source,native)
        slides=[]; current={"title":"Presentation ORIA", "bullets":[]}
        for block in blocks:
            if block['type'] in {'h1','h2'}:
                if current['bullets']:slides.append(current)
                current={'title':block.get('text','Diapositive'), 'bullets':[]}
            elif block['type'] in {'bullet','number'}:
                current['bullets'].extend(block.get('items',[]))
            elif block['type']=='p':
                if block.get('text'):current['bullets'].append(block['text'])
        if current['bullets'] or not slides:slides.append(current)
        if len(slides)>50:raise ArtifactError('Presentation trop longue (50 slides maximum).')
        return slides

    def _contract_pptx(self, kind: str | None, source: str | None, native: Any) -> bytes:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
            from pptx.enum.shapes import MSO_SHAPE
        except ImportError as error:
            raise ArtifactDependencyError('PPTX : installer python-pptx.') from error
        slides=self._contract_slides(kind,source,native)
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        navy=RGBColor(24,51,93); accent=RGBColor(38,102,200)
        for index, spec in enumerate(slides,1):
            unsupported = set(spec) - {'title','subtitle','bullets','content','paragraphs',
                                       'table','elements','chart'}
            if unsupported:
                raise ArtifactError('Champ de diapositive non pris en charge : '
                                    + ', '.join(sorted(unsupported)))
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(250,252,255)
            bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.14))
            bar.fill.solid();bar.fill.fore_color.rgb=accent;bar.line.fill.background()
            title=str(spec.get('title') or f'Diapositive {index}')
            if len(title)>180:raise ArtifactError('Titre de diapositive trop long.')
            tx=slide.shapes.add_textbox(Inches(.75), Inches(.55), Inches(11.9), Inches(1.0))
            frame=tx.text_frame;frame.word_wrap=True
            p=frame.paragraphs[0];p.text=title;p.font.size=Pt(31 if len(title)<65 else 25)
            p.font.bold=True;p.font.color.rgb=navy
            subtitle=str(spec.get('subtitle') or '')
            start_y=1.72
            if subtitle:
                sub=slide.shapes.add_textbox(Inches(.85),Inches(1.55),Inches(11.7),Inches(.6))
                para=sub.text_frame.paragraphs[0];para.text=subtitle[:250];para.font.size=Pt(16)
                para.font.color.rgb=accent;start_y=2.25
            items=spec.get('bullets',spec.get('content',spec.get('paragraphs',[])))
            if isinstance(items,str):items=[items]
            if not isinstance(items,list):items=[]
            if len(items)>9 or sum(len(str(item)) for item in items)>1300:
                raise ArtifactError('Diapositive trop dense : repartir sur plusieurs diapositives.')
            table=spec.get('table')
            elements=spec.get('elements',[])
            if not isinstance(elements,list) or len(elements)>30:
                raise ArtifactError('Elements de diapositive invalides.')
            if spec.get('chart') is not None:
                if not isinstance(spec['chart'],dict):
                    raise ArtifactError('Graphique de diapositive invalide.')
                elements=[*elements, {**spec['chart'], 'type':'chart'}]
            if elements and (table is not None or items):
                raise ArtifactError('Ne pas melanger composition explicite et disposition automatique.')
            if elements:
                self._draw_contract_slide_elements(slide,elements)
            elif table is not None:
                if not isinstance(table,(list,dict)):
                    raise ArtifactError('Tableau de diapositive invalide.')
                rows=self._normalize_rows(table)
                if len(rows)>12 or max((len(row) for row in rows),default=0)>7:
                    raise ArtifactError('Tableau de diapositive trop large ou long.')
                if rows:
                    shape=slide.shapes.add_table(len(rows),max(len(row) for row in rows),
                            Inches(.85), Inches(start_y), Inches(11.7), Inches(4.45))
                    for r,row in enumerate(rows):
                        for c,val in enumerate(row):
                            cell=shape.table.cell(r,c);cell.text=str(val if val is not None else '')
                            cell.fill.solid();cell.fill.fore_color.rgb=(navy if r==0 else RGBColor(240,245,253))
                            for para in cell.text_frame.paragraphs:
                                for run in para.runs:
                                    run.font.size=Pt(13);run.font.color.rgb=(RGBColor(255,255,255) if r==0 else navy)
                    if items:raise ArtifactError('Diapositive avec tableau et liste : agencement non defini.')
            elif items:
                tb=slide.shapes.add_textbox(Inches(1),Inches(start_y),Inches(11.35),Inches(4.6))
                tf=tb.text_frame;tf.word_wrap=True;tf.clear()
                font_size=Pt(22 if len(items)<=5 else (19 if len(items)<=7 else 16))
                for rowno,item in enumerate(items):
                    para=tf.paragraphs[0] if rowno==0 else tf.add_paragraph()
                    para.text='•  '+str(item);para.font.size=font_size
                    para.font.color.rgb=navy;para.space_after=Pt(14 if len(items)<=6 else 8)
            footer=slide.shapes.add_textbox(Inches(11.65),Inches(7.05),Inches(.9),Inches(.25))
            f=footer.text_frame.paragraphs[0];f.text=f'{index}/{len(slides)}'
            f.font.size=Pt(10);f.font.color.rgb=RGBColor(110,120,138)
        output=io.BytesIO();prs.save(output);return output.getvalue()

    def _draw_contract_slide_elements(self, slide: Any, elements: list[dict[str, Any]]) -> None:
        """Objets PowerPoint natifs et modifiables, positions en pouces."""
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        def pos(v: Any, default: float, bound: float) -> float:
            value = default if v is None else v
            if isinstance(value, bool) or not isinstance(value,(int,float)) or not 0 <= float(value) <= bound:
                raise ArtifactError('Position ou dimension de slide invalide.')
            return float(value)
        def color(value: Any, default: str = '234269') -> Any:
            v = str(value or default).strip().lstrip('#')
            if not re.fullmatch(r'[0-9A-Fa-f]{6}',v):
                raise ArtifactError('Couleur de slide invalide.')
            return RGBColor.from_string(v.upper())
        for el in elements:
            if not isinstance(el,dict):raise ArtifactError('Objet de slide non structure.')
            typ=str(el.get('type','')).lower()
            x=pos(el.get('x'),.9,13.2); y=pos(el.get('y'),2,7.4)
            w=pos(el.get('width'),5,13.2); h=pos(el.get('height'),2,7.4)
            if w<.1 or h<.1 or x+w>13.35 or y+h>7.5:
                raise ArtifactError('Objet PowerPoint hors de la diapositive.')
            if typ=='text':
                text=str(el.get('text') or '')
                if len(text)>1300:raise ArtifactError('Bloc texte de slide trop long.')
                shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
                tf=shape.text_frame;tf.word_wrap=True
                para=tf.paragraphs[0];para.text=text
                para.font.size=Pt(max(8,min(48,pos(el.get('font_size'),19,48))))
                para.font.bold=bool(el.get('bold',False));para.font.color.rgb=color(el.get('color'))
            elif typ in {'rectangle','rect'}:
                shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y), Inches(w), Inches(h))
                shape.fill.solid();shape.fill.fore_color.rgb=color(el.get('fill'),'EBF2FE')
                shape.line.fill.background()
            elif typ=='image':
                _,raw=self._image_data_uri(str(el.get('src') or ''))
                slide.shapes.add_picture(io.BytesIO(raw), Inches(x),Inches(y), width=Inches(w),height=Inches(h))
            elif typ=='chart':
                kind=str(el.get('chart_type','bar')).lower()
                chart_type={'bar':XL_CHART_TYPE.COLUMN_CLUSTERED,
                            'line':XL_CHART_TYPE.LINE,
                            'pie':XL_CHART_TYPE.PIE}.get(kind)
                if chart_type is None:raise ArtifactError('Graphique PPTX non pris en charge.')
                cats=el.get('categories')
                series=el.get('series')
                if (not isinstance(cats,list) or not cats or len(cats)>40
                        or not isinstance(series,list) or not series or len(series)>8):
                    raise ArtifactError('Donnees de graphique PPTX invalides.')
                data=CategoryChartData();data.categories=[str(c) for c in cats]
                for entry in series:
                    if (not isinstance(entry,dict) or not isinstance(entry.get('values'),list)
                            or len(entry['values'])!=len(cats)
                            or any(isinstance(v,bool) or not isinstance(v,(int,float)) for v in entry['values'])):
                        raise ArtifactError('Serie de graphique PPTX invalide.')
                    data.add_series(str(entry.get('name') or 'Serie'), entry['values'])
                shape=slide.shapes.add_chart(chart_type, Inches(x),Inches(y),Inches(w),Inches(h),data)
                shape.chart.has_legend=len(series)>1
            else:
                raise ArtifactError(f'Type d objet PowerPoint non pris en charge : {typ}.')

    def _contract_sheets(self, kind: str | None, source: str | None,
                         native: Any) -> list[dict[str, Any]]:
        if isinstance(native, dict) and isinstance(native.get('sheets'), list):
            sheets = native['sheets']
            if not sheets or len(sheets)>20 or any(not isinstance(s,dict) for s in sheets):
                raise ArtifactError('Liste de feuilles de calcul invalide.')
            return sheets
        if native is not None:
            return [{"name":"Donnees", "rows":self._normalize_rows(native)}]
        if kind!='html' or not source:
            raise ArtifactError('Tableur sans source tabulaire exploitable.')
        from bs4 import BeautifulSoup
        soup=BeautifulSoup(source,'html.parser')
        tables=soup.find_all('table')
        if not tables:raise ArtifactError('Aucun tableau HTML pour le tableur.')
        result=[]
        for idx,table in enumerate(tables[:20],1):
            rows=[]
            for tr in table.find_all('tr'):
                values=[]
                for cell in tr.find_all(['td','th'],recursive=False):
                    text=cell.get_text(' ',strip=True)
                    if cell.get('data-type')=='number' and re.fullmatch(r'-?\d+(?:\.\d+)?',text):
                        values.append(float(text) if '.' in text else int(text))
                    else:values.append(text)
                if values: rows.append(values)
            result.append({'name':table.get('id') or f'Tableau {idx}','rows':rows})
        return result

    def _contract_xlsx(self, kind: str | None, source: str | None, native: Any) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        except ImportError as error:
            raise ArtifactDependencyError('XLSX : installer openpyxl.') from error
        sheets=self._contract_sheets(kind,source,native)
        book=Workbook();book.remove(book.active)
        used=set()
        for index,s in enumerate(sheets,1):
            name=self._safe_sheet_name(s.get('name') or f'Feuille {index}')
            if name.casefold() in used:raise ArtifactError('Nom de feuille en double.')
            used.add(name.casefold()); ws=book.create_sheet(name)
            rows=s.get('rows',s.get('data',[]))
            if isinstance(rows,list) and all(isinstance(r,dict) for r in rows):
                rows=self._normalize_rows(rows)
            if not isinstance(rows,list) or len(rows)>10000:
                raise ArtifactError('Feuille trop volumineuse ou sans lignes valides.')
            for r,row in enumerate(rows,1):
                if not isinstance(row,(list,tuple)) or len(row)>100:
                    raise ArtifactError('Ligne de tableur non valide ou trop large.')
                for c,value in enumerate(row,1):
                    fmt=None
                    if isinstance(value,dict):
                        fmt=value.get('number_format')
                        if 'formula' in value:
                            formula=value['formula']
                            if not isinstance(formula,str) or not formula.startswith('=') or len(formula)>800:
                                raise ArtifactError('Formule de cellule invalide.')
                            if re.search(r'(?i)\b(?:HYPERLINK|WEBSERVICE|RTD|DDE|CALL|REGISTER\.ID|FILTERXML)\s*\(|https?://|file:|\[[^\]]+\]',formula):
                                raise ArtifactError('Formule avec acces externe non autorisee.')
                            value=formula
                        else:value=value.get('value')
                    elif isinstance(value,str) and value.startswith('='):
                        # La cellule de texte ne devient pas une formule implicite.
                        value=value
                    if not isinstance(value,(str,int,float,bool,type(None))):
                        raise ArtifactError('Type de valeur de cellule non pris en charge.')
                    cell=ws.cell(r,c)
                    if isinstance(value,str) and value.startswith('=') and not (isinstance(row[c-1],dict) and 'formula' in row[c-1]):
                        cell.value="'"+value
                    else: cell.value=value
                    if fmt:
                        if not isinstance(fmt,str) or len(fmt)>80:raise ArtifactError('Format numerique invalide.')
                        cell.number_format=fmt
                    cell.alignment=Alignment(vertical='center',wrap_text=True)
                    if r==1:
                        cell.font=Font(name='Aptos',size=11,color='FFFFFF',bold=True)
                        cell.fill=PatternFill('solid',fgColor='18335D')
                    else:
                        cell.font=Font(name='Aptos',size=10,color='22334B')
                        if r%2==0:cell.fill=PatternFill('solid',fgColor='EDF3FC')
            if rows:
                ws.row_dimensions[1].height=25
                for col in range(1,min(len(rows[0]),100)+1):
                    width=min(55,max(13,max((len(str((r[col-1] if col<=len(r) else '') or '')) for r in rows[:150]),default=8)+3))
                    ws.column_dimensions[get_column_letter(col)].width=width
            ws.freeze_panes=str(s.get('freeze_panes') or ('A2' if rows else 'A1'))
            if s.get('auto_filter',True) and rows:
                ws.auto_filter.ref=ws.dimensions
            merged=s.get('merged_cells',[])
            if not isinstance(merged,list) or len(merged)>100:raise ArtifactError('Plages fusionnees invalides.')
            for merge in merged:
                if not isinstance(merge,str) or not re.fullmatch(r'[A-Z]{1,3}\d+:[A-Z]{1,3}\d+',merge):
                    raise ArtifactError('Plage de fusion invalide.')
                ws.merge_cells(merge)
            charts=s.get('charts',[])
            if not isinstance(charts,list) or len(charts)>10:raise ArtifactError('Graphiques invalides.')
            for chart in charts:
                if not isinstance(chart,dict):raise ArtifactError('Definition de graphique invalide.')
                typ=str(chart.get('type','bar')).lower()
                ch={'bar':BarChart,'line':LineChart,'pie':PieChart}.get(typ)
                if ch is None:raise ArtifactError('Type de graphique non supporte.')
                obj=ch();obj.title=str(chart.get('title','Graphique'))[:120]
                try:
                    min_col=int(chart.get('min_col',2));max_col=int(chart.get('max_col',min_col))
                    min_row=int(chart.get('min_row',1));max_row=int(chart.get('max_row',len(rows)))
                    if not (1<=min_col<=max_col<=100 and 2<=max_row<=10000 and 1<=min_row<=max_row):
                        raise ValueError('Plage graphique invalide')
                    data=Reference(ws,min_col=min_col,max_col=max_col,min_row=min_row,max_row=max_row)
                    obj.add_data(data,titles_from_data=(min_row==1))
                    obj.set_categories(Reference(ws,min_col=int(chart.get('category_col',1)),min_row=max(2,min_row+1),max_row=max_row))
                    ws.add_chart(obj,str(chart.get('anchor','H2')))
                except (TypeError,ValueError) as error:
                    raise ArtifactError('Parametres de graphique invalides.') from error
        buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()

    def _contract_ods(self, kind: str | None, source: str | None, native: Any) -> bytes:
        try:
            from odf.opendocument import OpenDocumentSpreadsheet
            from odf.table import Table, TableCell, TableRow
            from odf.text import P
        except ImportError as error:
            raise ArtifactDependencyError('ODS : installer odfpy.') from error
        doc=OpenDocumentSpreadsheet()
        for index,s in enumerate(self._contract_sheets(kind,source,native),1):
            table=Table(name=self._safe_sheet_name(s.get('name') or f'Feuille {index}'))
            rows=s.get('rows',s.get('data',[]))
            if isinstance(rows,list) and all(isinstance(r,dict) for r in rows):
                rows=self._normalize_rows(rows)
            if not isinstance(rows,list) or len(rows)>10000:raise ArtifactError('Lignes ODS invalides.')
            for row in rows:
                if not isinstance(row,(list,tuple)) or len(row)>100:raise ArtifactError('Ligne ODS invalide.')
                tr=TableRow()
                for value in row:
                    formula=None
                    if isinstance(value,dict):
                        formula=value.get('formula');value=value.get('value')
                    if formula is not None:
                        if not isinstance(formula,str) or not formula.startswith('='):
                            raise ArtifactError('Formule ODS invalide.')
                        cell=TableCell(valuetype='float',value=0,formula='of:'+formula)
                    elif isinstance(value,bool):cell=TableCell(valuetype='boolean',booleanvalue=str(value).lower())
                    elif isinstance(value,(int,float)) and not isinstance(value,bool):
                        cell=TableCell(valuetype='float',value=value)
                    else:cell=TableCell(valuetype='string')
                    cell.addElement(P(text='' if value is None else str(value)))
                    tr.addElement(cell)
                table.addElement(tr)
            doc.spreadsheet.addElement(table)
        out=io.BytesIO();doc.write(out);return out.getvalue()

    def _contract_delimited(self, fmt: str, kind: str | None,
                            source: str | None, native: Any) -> bytes:
        if native is None:
            sheets=self._contract_sheets(kind,source,native)
            if len(sheets)!=1:raise ArtifactError('CSV/TSV ne peut contenir plusieurs tableaux.')
            rows=sheets[0]['rows']
        else:
            if isinstance(native,dict) and 'sheets' in native:
                sheets=self._contract_sheets(kind,source,native)
                if len(sheets)!=1:raise ArtifactError('CSV/TSV ne contient qu une seule feuille.')
                rows=sheets[0].get('rows',[])
            else:rows=self._normalize_rows(native)
        safe=[]
        for row in rows:
            values=[]
            for val in row:
                if isinstance(val,dict):val=val.get('value', '')
                value='' if val is None else str(val)
                # CSV/TSV n'a pas de type FORMULE ; empecher l'injection Excel.
                if value.lstrip().startswith(('=','+','-','@')) and not re.fullmatch(r'-?\d+(?:\.\d+)?',value):
                    value="'"+value
                values.append(value)
            safe.append(values)
        return self._create_delimited(safe,delimiter=',' if fmt=='csv' else '\t')

    def _contract_sqlite(self, native: Any) -> bytes:
        tables=self._normalize_database_tables(native)
        if len(tables)>25:raise ArtifactError('SQLite : trop de tables.')
        conn=sqlite3.connect(':memory:')
        try:
            names=set()
            for table in tables:
                name=self._safe_sql_identifier(table['name'])
                if name.casefold() in names:raise ArtifactError('Table SQLite dupliquee.')
                names.add(name.casefold())
                rows=table.get('rows',[])
                raw_table=next((t for t in native.get('tables',[]) if isinstance(t,dict) and t.get('name')==table['name']),None) if isinstance(native,dict) else None
                columns=(raw_table or {}).get('columns')
                if columns is not None:
                    if not isinstance(columns,list) or not columns:raise ArtifactError('Colonnes SQLite invalides.')
                    cols=[]
                    for c in columns:
                        if not isinstance(c,dict) or not isinstance(c.get('name'),str):
                            raise ArtifactError('Definition de colonne SQLite invalide.')
                        dtype=str(c.get('type','TEXT')).upper()
                        if dtype not in {'TEXT','INTEGER','REAL','NUMERIC'}:raise ArtifactError('Type SQLite interdit.')
                        cols.append((self._safe_sql_identifier(c['name']),dtype))
                elif rows:
                    cols=[(self._safe_sql_identifier(k),'TEXT') for k in rows[0]]
                else: cols=[('_empty','TEXT')]
                if len(cols)>100 or len(set(k.casefold() for k,t in cols))!=len(cols):
                    raise ArtifactError('Colonnes SQLite en double ou trop nombreuses.')
                quoted=', '.join(f'"{col}" {typ}' for col,typ in cols)
                conn.execute(f'CREATE TABLE "{name}" ({quoted})')
                if len(rows)>50000:raise ArtifactError('Table SQLite trop volumineuse.')
                if rows:
                    keys=[k for k,_ in cols]
                    sql=f'INSERT INTO "{name}" VALUES ({",".join("?" for _ in cols)})'
                    vals=[]
                    for row in rows:
                        if not isinstance(row,dict):raise ArtifactError('Ligne SQLite non structuree.')
                        values=[]
                        for col,typ in cols:
                            v=row.get(col)
                            if v is not None and not isinstance(v,(str,int,float,bool)):
                                raise ArtifactError('Valeur SQLite incompatible.')
                            values.append(v)
                        vals.append(values)
                    conn.executemany(sql,vals)
            conn.commit()
            return bytes(conn.serialize())
        finally:conn.close()

    # ============================================================
    # NORMALISATEURS DE CONTENU
    # ============================================================

    def _normalize_rows(
        self,
        content: Any,
    ) -> list[list[Any]]:
        if isinstance(content, dict):
            if "rows" in content:
                content = content["rows"]
            elif "data" in content:
                content = content["data"]
            else:
                content = [content]

        if content is None:
            return []

        if isinstance(content, str):
            return [[content]]

        if not isinstance(content, (list, tuple)):
            return [[content]]

        if not content:
            return []

        if all(isinstance(item, dict) for item in content):
            headers: list[str] = []
            seen: set[str] = set()

            for row in content:
                for key in row.keys():
                    key_text = str(key)
                    if key_text not in seen:
                        seen.add(key_text)
                        headers.append(key_text)

            rows: list[list[Any]] = [headers]

            for row in content:
                rows.append(
                    [
                        row.get(header)
                        for header in headers
                    ]
                )

            return rows

        normalized: list[list[Any]] = []

        for item in content:
            if isinstance(item, (list, tuple)):
                normalized.append(list(item))
            else:
                normalized.append([item])

        return normalized

    def _normalize_sheets(
        self,
        content: Any,
    ) -> list[dict[str, Any]]:
        if isinstance(content, dict) and isinstance(
            content.get("sheets"),
            list,
        ):
            result = []

            for index, sheet in enumerate(
                content["sheets"],
                start=1,
            ):
                if isinstance(sheet, dict):
                    rows = self._normalize_rows(
                        sheet.get("rows", sheet.get("data", []))
                    )
                    result.append(
                        {
                            "name": sheet.get(
                                "name",
                                f"Feuille {index}",
                            ),
                            "rows": rows,
                        }
                    )

            return result or [
                {
                    "name": "Feuille 1",
                    "rows": [],
                }
            ]

        return [
            {
                "name": "Feuille 1",
                "rows": self._normalize_rows(content),
            }
        ]

    def _normalize_document_blocks(
        self,
        content: Any,
    ) -> list[dict[str, Any]]:
        if isinstance(content, dict) and isinstance(
            content.get("blocks"),
            list,
        ):
            normalized = []
            for block in content["blocks"]:
                if isinstance(block, dict):
                    normalized.append(
                        {
                            "type": str(
                                block.get("type") or "p"
                            ).lower(),
                            "text": str(
                                block.get("text") or ""
                            ),
                            "items": [
                                str(item)
                                for item in block.get("items", [])
                            ],
                            "rows": self._normalize_rows(
                                block.get("rows", [])
                            ),
                        }
                    )
            return normalized

        if isinstance(content, (list, tuple, dict)):
            if isinstance(content, dict):
                text = json.dumps(
                    content,
                    ensure_ascii=False,
                    indent=2,
                    default=str,
                )
            else:
                text = "\n".join(str(item) for item in content)
        else:
            text = "" if content is None else str(content)

        blocks: list[dict[str, Any]] = []
        bullet_buffer: list[str] = []

        def flush_bullets() -> None:
            nonlocal bullet_buffer
            if bullet_buffer:
                blocks.append(
                    {
                        "type": "bullet",
                        "items": bullet_buffer,
                    }
                )
                bullet_buffer = []

        paragraph_buffer: list[str] = []

        def flush_paragraph() -> None:
            nonlocal paragraph_buffer
            if paragraph_buffer:
                blocks.append(
                    {
                        "type": "p",
                        "text": "\n".join(paragraph_buffer).strip(),
                    }
                )
                paragraph_buffer = []

        for raw_line in text.splitlines():
            line = raw_line.rstrip()

            if not line.strip():
                flush_paragraph()
                flush_bullets()
                continue

            match = re.match(r"^(#{1,3})\s+(.+)$", line)
            if match:
                flush_paragraph()
                flush_bullets()
                blocks.append(
                    {
                        "type": f"h{len(match.group(1))}",
                        "text": match.group(2).strip(),
                    }
                )
                continue

            bullet = re.match(r"^\s*[-*•]\s+(.+)$", line)
            if bullet:
                flush_paragraph()
                bullet_buffer.append(
                    bullet.group(1).strip()
                )
                continue

            flush_bullets()
            paragraph_buffer.append(line)

        flush_paragraph()
        flush_bullets()

        return blocks or [{"type": "p", "text": ""}]

    def _document_plain_text(
        self,
        content: Any,
    ) -> str:
        parts: list[str] = []

        for block in self._normalize_document_blocks(content):
            kind = block["type"]

            if kind == "bullet":
                parts.extend(
                    f"• {item}"
                    for item in block.get("items", [])
                )
            elif kind == "table":
                parts.extend(
                    " | ".join(
                        "" if value is None else str(value)
                        for value in row
                    )
                    for row in block.get("rows", [])
                )
            else:
                parts.append(block.get("text", ""))

        return "\n\n".join(parts)

    def _normalize_slides(
        self,
        content: Any,
        *,
        title: str | None,
    ) -> list[dict[str, Any]]:
        if isinstance(content, dict) and isinstance(
            content.get("slides"),
            list,
        ):
            slides = []

            for index, slide in enumerate(
                content["slides"],
                start=1,
            ):
                if isinstance(slide, dict):
                    bullets = slide.get("bullets", slide.get("content", []))

                    if isinstance(bullets, str):
                        bullets = [
                            line.strip("-• ").strip()
                            for line in bullets.splitlines()
                            if line.strip()
                        ]

                    if not isinstance(bullets, list):
                        bullets = [bullets]

                    slides.append(
                        {
                            "title": str(
                                slide.get("title")
                                or f"Diapositive {index}"
                            ),
                            "bullets": [
                                str(item)
                                for item in bullets
                                if item is not None
                            ],
                        }
                    )

            if slides:
                return slides

        blocks = self._normalize_document_blocks(content)
        slides: list[dict[str, Any]] = []
        current = {
            "title": title or "Présentation Oria",
            "bullets": [],
        }

        for block in blocks:
            if block["type"] in {"h1", "h2"}:
                if current["bullets"] or slides:
                    slides.append(current)
                current = {
                    "title": block.get("text") or "Diapositive",
                    "bullets": [],
                }
            elif block["type"] == "bullet":
                current["bullets"].extend(block.get("items", []))
            elif block["type"] == "p":
                paragraph = block.get("text", "").strip()
                if paragraph:
                    current["bullets"].append(paragraph)

        if current["title"] or current["bullets"]:
            slides.append(current)

        return slides or [
            {
                "title": title or "Présentation Oria",
                "bullets": [],
            }
        ]

    def _normalize_epub_chapters(
        self,
        content: Any,
    ) -> list[dict[str, Any]]:
        if isinstance(content, dict) and isinstance(
            content.get("chapters"),
            list,
        ):
            chapters = []

            for index, chapter in enumerate(
                content["chapters"],
                start=1,
            ):
                if isinstance(chapter, dict):
                    raw = chapter.get(
                        "content",
                        chapter.get("text", ""),
                    )
                    paragraphs = [
                        value.strip()
                        for value in str(raw).split("\n\n")
                        if value.strip()
                    ]
                    chapters.append(
                        {
                            "title": str(
                                chapter.get("title")
                                or f"Chapitre {index}"
                            ),
                            "paragraphs": paragraphs or [""],
                        }
                    )

            if chapters:
                return chapters

        return [
            {
                "title": "Document",
                "paragraphs": [
                    block.get("text", "")
                    for block in self._normalize_document_blocks(content)
                    if block.get("text", "").strip()
                ] or [""],
            }
        ]

    def _normalize_database_tables(
        self,
        content: Any,
    ) -> list[dict[str, Any]]:
        if isinstance(content, dict) and isinstance(
            content.get("tables"),
            list,
        ):
            tables = []

            for index, table in enumerate(
                content["tables"],
                start=1,
            ):
                if not isinstance(table, dict):
                    continue

                rows = table.get("rows", table.get("data", []))

                if isinstance(rows, dict):
                    rows = [rows]

                if not isinstance(rows, list):
                    rows = []

                normalized_rows = [
                    row
                    for row in rows
                    if isinstance(row, dict)
                ]

                tables.append(
                    {
                        "name": str(
                            table.get("name")
                            or f"table_{index}"
                        ),
                        "rows": normalized_rows,
                    }
                )

            return tables or [
                {
                    "name": "data",
                    "rows": [],
                }
            ]

        if isinstance(content, list) and all(
            isinstance(row, dict)
            for row in content
        ):
            return [
                {
                    "name": "data",
                    "rows": content,
                }
            ]

        if isinstance(content, dict):
            return [
                {
                    "name": "data",
                    "rows": [content],
                }
            ]

        raise ArtifactError(
            "SQLite attend une liste de dictionnaires ou "
            "{'tables': [{'name': ..., 'rows': [...]}]}."
        )

    # ============================================================
    # HELPERS
    # ============================================================

    @staticmethod
    def _safe_sheet_name(
        value: Any,
    ) -> str:
        name = str(value or "Feuille").strip()
        name = re.sub(r"[\[\]:*?/\\]+", "-", name)
        return (name[:31] or "Feuille").strip()

    @staticmethod
    def _safe_sql_identifier(
        value: Any,
    ) -> str:
        name = str(value or "data").strip()
        name = re.sub(r"[^0-9A-Za-z_ -]+", "_", name)
        name = name.replace('"', "_")
        return name[:80] or "data"

    @staticmethod
    def _safe_archive_name(
        value: str,
    ) -> str:
        # Autorise les sous-dossiers mais interdit traversée et chemins absolus.
        raw = value.replace("\\", "/").strip()
        path = PurePosixPath(raw)

        parts = [
            part
            for part in path.parts
            if part not in {"", ".", "..", "/"}
        ]

        if not parts:
            return "fichier.txt"

        cleaned = "/".join(parts)
        cleaned = re.sub(r"[\x00-\x1f<>:\"|?*]+", "-", cleaned)
        return cleaned[:240]

    @staticmethod
    def _validate_project_archive_name(
        original_name: str,
        archive_name: str,
    ) -> None:
        """Refuse les noms dangereux ou altérés pour les fichiers source ZIP."""
        raw = original_name.replace("\\", "/").strip()
        parts = raw.split("/")
        reserved = {"CON", "PRN", "AUX", "NUL"}
        reserved.update(f"COM{i}" for i in range(1, 10))
        reserved.update(f"LPT{i}" for i in range(1, 10))
        windows_incompatible = any(
            part.split(".", 1)[0].upper() in reserved
            or part.endswith((".", " "))
            for part in parts
        )
        if (
            not raw
            or raw.startswith("/")
            or re.match(r"^[A-Za-z]:", raw)
            or any(part in {"", ".", ".."} for part in parts)
            or windows_incompatible
            or len(raw) > 240
            or raw != archive_name
        ):
            raise ArtifactError(
                f"Chemin de fichier source ZIP invalide : {original_name!r}."
            )

    @staticmethod
    def _deduplicate_filename(
        filename: str,
        used_names: set[str],
    ) -> str:
        path = PurePosixPath(filename)
        parent = "" if str(path.parent) == "." else f"{path.parent}/"
        stem = path.stem
        suffix = path.suffix

        index = 2
        candidate = filename

        while candidate in used_names:
            candidate = f"{parent}{stem}-{index}{suffix}"
            index += 1

        return candidate
